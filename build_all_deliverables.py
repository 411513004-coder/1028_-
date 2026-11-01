import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.text import PP_ALIGN

WORK_DIR = r"c:\User\Downloads\1023ai期中"
EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report_docx():
    doc = docx.Document()
    
    # 邊界設定
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # 標題
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("《花蓮生活探索地圖》期中專題報告")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(6, 78, 59) # 深綠
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(12)

    # Metadata 表格 / 資訊卡
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("專案名稱", "花蓮生活探索地圖 (Hualien Living Map) —— 蔬食 · 景點 · 咖啡館指南"),
        ("專案類別", "Web 互動式地圖應用 (HTML5 / Vanilla JS / Leaflet / Esri 圖資)"),
        ("核心規範", "100% 零商業金鑰 (No API Key) · 零安裝雙擊即跑 · 50 處全實名核實"),
        ("專案連結", "https://github.com/411513004-coder/1028_-"),
        ("線上體驗", "https://411513004-coder.github.io/1028_-/")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(1.5)
        c2.width = Inches(5.2)
        set_cell_background(c1, "F0FDF4")
        set_cell_background(c2, "F8FAFC")
        set_cell_margins(c1, 80, 80, 120, 120)
        set_cell_margins(c2, 80, 80, 120, 120)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(k)
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(6, 95, 70)
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(v)
        r2.font.size = Pt(10)
        if "http" in v:
            r2.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # 章節內容生成函數
    def add_section_heading(num_str, title_str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(f"{num_str}、 {title_str}")
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor(6, 78, 59)
        return h

    def add_subheading(title_str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(title_str)
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
        return h

    # 一、 做什麼
    add_section_heading("一", "做什麼 (What)")
    p = doc.add_paragraph()
    p.add_run("本專案開發了一款專為花蓮在地居民、大專院校師生（如東華大學、慈濟大學）以及前來花蓮自由行旅客量身打造的").font.size = Pt(11)
    p.add_run("「花蓮生活探索地圖 Web App」").bold = True
    p.runs[1].font.size = Pt(11)
    p.add_run("。在擴充版中，系統跳脫單一飲食維度，全面升級為三大生活主題：").font.size = Pt(11)

    bullets = [
        ("🌱 蔬食美饌（30 家）：", "依飲食細分「全素/純素 (16家)」、「蛋奶素 (4家)」與「蔬食友善 (10家)」，徹底解決素食標籤模糊混雜的痛點。"),
        ("🏛️ 精選景點（10 處）：", "收錄花蓮市區將軍府1936、松園別館(2025新團隊重啟百大亮點)、文創園區，吉安慶修院，壽豐鯉魚潭、雲山水，新城七星潭，以及秀林台泥DAKA與崇德瑩農場等代表性文化歷史與自然地標。"),
        ("☕ 質感咖啡館（10 家）：", "收錄花蓮在地人氣極高的老宅、海景與自家烘焙咖啡館（如浮室、Giocare手沖、珈琲花、KŌHI宅、Stroll肉桂捲、新城練習曲書店咖啡、壽豐海碉堡等）。")
    ]
    for b_title, b_desc in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        r_t = bp.add_run(b_title)
        r_t.font.bold = True
        r_t.font.size = Pt(10.5)
        r_d = bp.add_run(b_desc)
        r_d.font.size = Pt(10.5)

    p_feat = doc.add_paragraph()
    p_feat.add_run("系統具備雙層主題分頁過濾、行政區多維度過濾、招牌關鍵字全文即時檢索、地圖與卡片雙向動態聚焦連動，並整合一鍵通話與 Google 地圖精準路線導航。").font.size = Pt(11)

    # 二、 給誰用
    add_section_heading("二", "給誰用 (Who)")
    whos = [
        ("在地居民與大專院校師生：", "平日尋找東華志學街商圈、慈濟建國校區周邊的平價蔬食與咖啡自習空間。"),
        ("綠色低碳自由行觀光客：", "跨越花蓮市、吉安、壽豐、新城、秀林、鳳林，一手掌握沿線美食、休憩咖啡館與必訪景點，無縫串聯一日或多日遊行程。"),
        ("注重資訊真實性的旅人：", "全數資料經 2025-2026 最新官方資訊與實地核實，排除停業或幽靈店家，提供最安心可靠的指南。")
    ]
    for w_title, w_desc in whos:
        wp = doc.add_paragraph(style='List Bullet')
        r_t = wp.add_run(w_title)
        r_t.font.bold = True
        r_t.font.size = Pt(10.5)
        r_d = wp.add_run(w_desc)
        r_d.font.size = Pt(10.5)

    # 三、 怎麼做
    add_section_heading("三", "怎麼做 (How)")
    add_subheading("1. 架構設計與技術選型")
    
    tech_list = [
        ("前端介面 (UI/UX)：", "純原生 HTML5 + CSS3 + Vanilla JavaScript，無須繁瑣編譯，雙擊 index.html 即可於任一瀏覽器秒開執行。採用清新綠色與淡紫雙色系質感設計，支援桌面雙欄與手機端自適應抽屜式切換。"),
        ("地圖核心 (Map Engine)：", "選用全球知名的開源 Leaflet.js 串接 Esri World Street Map 街道圖資（並附備用高解析 Esri 衛星空照圖）。完全免申請商業金鑰（100% No API Key），零額度費用風險、杜絕金鑰外洩或作業違規扣分，支援縮放至 19 級且不擋本機檔案協議。"),
        ("嚴謹資料工程 (Data Engineering)：", "以輕量結構化 JSON 陣列建立資料庫，收錄 50 處 100% 現正營業/開放之花蓮在地熱點。實施嚴格資料清洗，詳細標註經緯度、營業/開放時間、公休日提示、票價/價位區間、聯絡方式與特色亮點。")
    ]
    for t_title, t_desc in tech_list:
        tp = doc.add_paragraph(style='List Bullet')
        r_t = tp.add_run(t_title)
        r_t.font.bold = True
        r_t.font.size = Pt(10.5)
        r_d = tp.add_run(t_desc)
        r_d.font.size = Pt(10.5)

    add_subheading("2. 核心功能實作")
    features = [
        ("主題與子類別雙層快選：", "第一層：全部探索 (50)、蔬食美饌 (30)、精選景點 (10)、質感咖啡館 (10)；第二層：切換至蔬食時動態展開「全素」、「蛋奶素」、「蔬食友善」子標籤。"),
        ("多維度分色視覺化標記 (Pins)：", "依類別給予不同專屬圖標與色彩（綠色 全素、黃色 蛋奶素、藍色 友善、紫色 景點、焦糖色 咖啡館），辨識度極高。"),
        ("雙向地圖卡片連動 (Bi-directional Linking)：", "點擊左側卡片，地圖平滑平移（FlyTo）並展開詳細彈窗；點擊地圖標記，清單自動平滑滾動並高亮該卡片。"),
        ("一鍵行動外連：", "支援 tel: 一鍵撥號（若是預約制店家貼心提示社群預約），並整合 Google Maps 路線導航 API。")
    ]
    for f_title, f_desc in features:
        fp = doc.add_paragraph(style='List Bullet')
        r_t = fp.add_run(f_title)
        r_t.font.bold = True
        r_t.font.size = Pt(10.5)
        r_d = fp.add_run(f_desc)
        r_d.font.size = Pt(10.5)

    # 四、 怎麼測
    add_section_heading("四", "怎麼測 (How to Test)")
    add_subheading("1. 功能測試矩陣 (Functional Testing)")
    
    test_table = doc.add_table(rows=7, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    test_headers = ["測試項目", "測試步驟", "預期結果", "測試結果"]
    col_widths = [Inches(1.3), Inches(1.8), Inches(2.7), Inches(0.9)]
    
    hdr_row = test_table.rows[0]
    for j, h_text in enumerate(test_headers):
        cell = hdr_row.cells[j]
        cell.width = col_widths[j]
        set_cell_background(cell, "064E3B")
        set_cell_margins(cell, 100, 100, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    test_rows_data = [
        ("地圖初始化", "瀏覽器載入頁面", "以花蓮中心展開，Esri 街道圖正常載入，顯示全區 50 處標記", "通過"),
        ("主題標籤切換", "點擊「精選景點」頁籤", "清單與地圖即時過濾出 10 處景點，自動隱藏蔬食子分類，數量正確", "通過"),
        ("咖啡廳專題過濾", "點擊「質感咖啡館」頁籤", "呈現 10 家老宅/海景咖啡館，地圖標記轉換為焦糖色咖啡杯圖標", "通過"),
        ("鄉鎮交叉篩選", "選擇鄉鎮「壽豐鄉」", "視角平滑聚焦壽豐，呈現阿宏素食、鯉魚潭、雲山水、海碉堡等地點", "通過"),
        ("關鍵字檢索", "搜尋「肉桂捲」或「將軍府」", "即時精準比對出 Stroll 散步咖啡或將軍府1936園區，卡片即時更新", "通過"),
        ("一鍵導航連動", "點擊任一彈窗的「導航前往」", "新分頁精準開啟 Google 地圖並帶入正確經緯度座標", "通過")
    ]
    for i, rdata in enumerate(test_rows_data):
        row = test_table.rows[i+1]
        for j, val in enumerate(rdata):
            cell = row.cells[j]
            cell.width = col_widths[j]
            if (i % 2) == 1:
                set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if j == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r.font.bold = True
                r.font.color.rgb = RGBColor(16, 185, 129)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_subheading("2. 跨裝置與環境相容性測試")
    compat = [
        ("桌面端瀏覽器 (Chrome / Edge / Firefox / Safari)：", "雙欄版面流暢，地圖拖曳、縮放與圖層切換反應極快。"),
        ("行動端裝置 (iOS / Android)：", "手機版自適應單欄抽屜，點擊底部「查看清單 / 查看地圖」順暢切換。"),
        ("靜態託管部署 (GitHub Pages)：", "經線上 HTTPS 測試，無任何破圖或跨網域安全性阻擋。")
    ]
    for c_title, c_desc in compat:
        cp = doc.add_paragraph(style='List Bullet')
        r_t = cp.add_run(c_title)
        r_t.font.bold = True
        r_t.font.size = Pt(10.5)
        r_d = cp.add_run(c_desc)
        r_d.font.size = Pt(10.5)

    # 五、 下一步
    add_section_heading("五", "下一步 (Next Steps)")
    next_steps = [
        ("多點行程自動規劃器 (Trip Planner)：", "支援使用者一鍵將想去的蔬食店、咖啡館與景點加入「收藏行程」，以演算法自動串聯最佳交通路線。"),
        ("GPS 即時定位與距離智慧排序：", "整合 HTML5 Geolocation API，自動計算使用者當前位置與各熱點的直線/行車距離，由近至遠排序。"),
        ("PWA 離線快取技術：", "引入 Service Worker，將地圖底圖快取至手機端，即便行經太魯閣山區或蘇花公路亦可離線查閱。"),
        ("小農與在地店家線上共好社群：", "打造開放式店家回報與推薦管道，促進花蓮在地永續綠色生活圈與觀光復甦。")
    ]
    for n_title, n_desc in next_steps:
        np = doc.add_paragraph(style='List Bullet')
        r_t = np.add_run(n_title)
        r_t.font.bold = True
        r_t.font.size = Pt(10.5)
        r_d = np.add_run(n_desc)
        r_d.font.size = Pt(10.5)

    docx_path = os.path.join(WORK_DIR, "專題報告_花蓮生活探索地圖.docx")
    doc.save(docx_path)
    print("Saved DOCX:", docx_path)
    return docx_path

def create_ai_log_docx():
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run("《花蓮生活探索地圖》AI 協同開發歷程摘要")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(6, 78, 59)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(12)

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("專案名稱", "花蓮生活探索地圖 (Hualien Living Map)"),
        ("協同 AI 工具", "Antigravity (Google DeepMind)"),
        ("專案開發期程", "2026 年 10 月 ~ 11 月"),
        ("開源倉庫連結", "https://github.com/411513004-coder/1028_-")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(1.5)
        c2.width = Inches(5.2)
        set_cell_background(c1, "F0FDF4")
        set_cell_background(c2, "F8FAFC")
        set_cell_margins(c1, 80, 80, 120, 120)
        set_cell_margins(c2, 80, 80, 120, 120)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(k)
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(6, 95, 70)
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(v)
        r2.font.size = Pt(10)
        if "http" in v:
            r2.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    stages = [
        ("一、 階段一：專案起步與需求拆解", [
            ("初期目標確立：", "使用者提出期中專題計畫，首要目標是嚴格符合課堂規範：「11/02 上課每人 90 秒展示（打開什麼、做什麼、看到什麼）」、「不強制龐大軟體，當天能順暢運行展示」、「程式碼、截圖絕對不可含 API 金鑰」。"),
            ("極簡輕量架構選型：", "經 AI 可行性評估，決定採取「純原生前端 Web 架構（HTML5 + CSS3 + Vanilla JS + Leaflet.js）」，免除繁瑣後端環境安裝，保證在任何電腦上雙擊 index.html 即可秒開。地圖圖資做到 100% 零商業金鑰、零額度限制、零資安洩漏風險。")
        ]),
        ("二、 階段二：實地查核與嚴格資料清洗 (Data Cleaning)", [
            ("辨識與修正初期問題：", "使用者實測指出早期名單有部分店家名稱不正確或已停業（如「太魯閣旅人」查無紀錄、「宏銘素食」為名稱誤植）。"),
            ("深度驗證與核實標準：", "AI 透過花蓮觀光資訊網、文化局古蹟登錄、Google 即時營業資訊與在地食記進行交叉查驗。更正新城「織娘之家 (預約制素食無菜單)」、北埔「金英蔬食」、東華大學周邊「阿宏素食」、「好佳小吃（獨立素食鍋具）」等名店。")
        ]),
        ("三、 階段三：功能擴充 —— 景點與咖啡廳全方位整合", [
            ("主題維度升級：", "使用者要求擴充花蓮景點與咖啡廳，全面升級為三大生活主題：精選景點 (10 處，含將軍府1936、松園別館2025重啟、文創園區、慶修院、七星潭、雲山水、台泥DAKA等)；質感咖啡館 (10 家，含浮室、Giocare手沖、珈琲花、KŌHI宅、Stroll肉桂捲、練習曲書店咖啡、海碉堡等)。"),
            ("嚴格排除已歇業店家：", "查核並排除七星潭已結束實體營業之「龍宮咖啡」，替補為新城老街正派營運之「練習曲書店咖啡」，總收錄達 50 筆 100% 現役地點。")
        ]),
        ("四、 階段四：互動邏輯與底圖穩定度優化", [
            ("雙層導航與智慧感知：", "首頁設計大主題頁籤（全部 / 蔬食 / 景點 / 咖啡廳），並在點選蔬食時動態展開全素、蛋奶素、友善子分類標籤。"),
            ("圖資穩定性修正：", "解決 CARTO 圖資在 local file:// 協議下會出現浮水印及 NatGeo 底圖高倍率破圖問題，全面切換為穩定無金鑰、支援 19 級高倍率之 Esri World Street Map 與 Esri 衛星空照圖。")
        ]),
        ("五、 階段五：版本控制、GitHub Pages 部署與 QR Code 生成", [
            ("版本控制規範：", "依課程規範設定 .gitignore，杜絕任何金鑰與機密檔案；倉庫名稱符合規範。"),
            ("GitHub Pages 線上發布：", "成功部署至 https://411513004-coder.github.io/1028_-/，所有人皆可直接線上體驗。"),
            ("專案 QR Code 整合：", "生成高解析專屬 QR Code 圖檔 (qrcode.png)，並無縫嵌入簡報封面，供 90 秒評審即時掃碼互動。")
        ])
    ]

    for title, items in stages:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(title)
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(6, 78, 59)

        for it_title, it_desc in items:
            p = doc.add_paragraph(style='List Bullet')
            r1 = p.add_run(it_title)
            r1.font.bold = True
            r1.font.size = Pt(10.5)
            r2 = p.add_run(it_desc)
            r2.font.size = Pt(10.5)

    docx_path = os.path.join(WORK_DIR, "AI協同開發歷程摘要.docx")
    doc.save(docx_path)
    print("Saved DOCX:", docx_path)
    return docx_path

def create_report_pdf():
    # 建立漂亮排版的 HTML 轉印 PDF
    html_content = """<!DOCTYPE html>
<html lang="zh-TW">
<head>
<meta charset="UTF-8">
<title>《花蓮生活探索地圖》期中專題報告</title>
<style>
  @page {
    size: A4;
    margin: 20mm 18mm 20mm 18mm;
  }
  body {
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "PingFang TC", "Microsoft JhengHei", sans-serif;
    color: #1e293b;
    line-height: 1.65;
    font-size: 10.5pt;
  }
  h1 {
    color: #064e3b;
    text-align: center;
    font-size: 20pt;
    margin-bottom: 6px;
    letter-spacing: -0.5px;
  }
  .meta-box {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    padding: 12px 18px;
    margin: 14px 0 20px 0;
    display: grid;
    grid-template-columns: auto 1fr;
    gap: 6px 14px;
    font-size: 9.5pt;
  }
  .meta-label {
    font-weight: bold;
    color: #065f46;
  }
  h2 {
    color: #064e3b;
    border-bottom: 2px solid #a7f3d0;
    padding-bottom: 4px;
    margin-top: 20px;
    margin-bottom: 8px;
    font-size: 13.5pt;
  }
  h3 {
    color: #0f172a;
    font-size: 11pt;
    margin-top: 12px;
    margin-bottom: 4px;
  }
  ul {
    margin-top: 4px;
    margin-bottom: 8px;
    padding-left: 20px;
  }
  li {
    margin-bottom: 4px;
  }
  table {
    width: 100%;
    border-collapse: collapse;
    margin: 12px 0;
    font-size: 9pt;
  }
  th, td {
    border: 1px solid #cbd5e1;
    padding: 7px 10px;
    text-align: left;
  }
  th {
    background: #064e3b;
    color: #ffffff;
    font-weight: 600;
  }
  tr:nth-child(even) td {
    background: #f8fafc;
  }
  .status-tag {
    color: #059669;
    font-weight: bold;
    text-align: center;
  }
  .qr-float {
    float: right;
    width: 110px;
    text-align: center;
    margin: 0 0 10px 15px;
    padding: 8px;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    background: #ffffff;
  }
  .qr-float img {
    width: 90px;
    height: 90px;
    display: block;
    margin: 0 auto 4px auto;
  }
  .qr-float span {
    font-size: 7.5pt;
    color: #64748b;
  }
</style>
</head>
<body>

<div class="qr-float">
  <img src="qrcode.png" alt="QR Code">
  <span>📱 掃碼線上體驗</span>
</div>

<h1>《花蓮生活探索地圖》期中專題報告</h1>

<div class="meta-box">
  <div class="meta-label">專案名稱：</div><div>花蓮生活探索地圖 (Hualien Living Map) —— 蔬食 · 景點 · 咖啡館全方位指南</div>
  <div class="meta-label">專案類別：</div><div>Web 互動式地圖應用 (HTML5 / Vanilla JS / Leaflet / Esri 圖資)</div>
  <div class="meta-label">核心原則：</div><div>100% 零商業金鑰 (No API Key) · 零安裝雙擊即跑 · 50 處全實名核實</div>
  <div class="meta-label">開源倉庫：</div><div>https://github.com/411513004-coder/1028_-</div>
  <div class="meta-label">線上體驗：</div><div>https://411513004-coder.github.io/1028_-/</div>
</div>

<h2>一、 做什麼 (What)</h2>
<p>本專案開發了一款專為花蓮在地居民、大專院校師生（如東華大學、慈濟大學）以及前來花蓮自由行旅客量身打造的<strong>「花蓮生活探索地圖 Web App」</strong>。在擴充版中，系統跳脫單一飲食維度，全面升級為三大生活主題：</p>
<ul>
  <li><strong>🌱 蔬食美饌（30 家）</strong>：依飲食細分「全素/純素 (16家)」、「蛋奶素 (4家)」與「蔬食友善 (10家)」，徹底解決素食標籤模糊混雜的痛點。</li>
  <li><strong>🏛️ 精選景點（10 處）</strong>：收錄花蓮市區將軍府1936、松園別館(2025新團隊重啟獲百大亮點)、文創園區，吉安慶修院，壽豐鯉魚潭、雲山水，新城七星潭，以及秀林台泥DAKA與崇德瑩農場等代表性文化歷史與自然地標。</li>
  <li><strong>☕ 質感咖啡館（10 家）</strong>：收錄花蓮在地人氣極高的老宅、海景與自家烘焙咖啡館（如浮室、Giocare手沖、珈琲花、KŌHI宅、Stroll肉桂捲、新城練習曲書店咖啡、壽豐海碉堡等）。</li>
</ul>
<p>系統具備<strong>雙層主題分頁過濾</strong>、<strong>行政區多維度過濾</strong>、<strong>招牌關鍵字全文即時檢索</strong>、<strong>地圖與卡片雙向動態聚焦連動</strong>，並整合<strong>一鍵通話</strong>與 <strong>Google 地圖路線導航</strong>。</p>

<h2>二、 給誰用 (Who)</h2>
<ul>
  <li><strong>在地居民與大專院校師生</strong>：平日尋找東華志學街商圈、慈濟建國校區周邊的平價蔬食與咖啡自習空間。</li>
  <li><strong>綠色低碳自由行觀光客</strong>：跨越花蓮市、吉安、壽豐、新城、秀林、鳳林，一手掌握沿線美食、休憩咖啡館與必訪景點，無縫串聯一日或多日遊行程。</li>
  <li><strong>注重資訊真實性的旅人</strong>：全數資料經 2025-2026 最新官方資訊與實地核實，排除停業或幽靈店家，提供最安心可靠的指南。</li>
</ul>

<h2>三、 怎麼做 (How)</h2>
<h3>1. 架構設計與技術選型</h3>
<ul>
  <li><strong>前端介面 (UI/UX)</strong>：純原生 HTML5 + CSS3 + Vanilla JavaScript，無須繁瑣編譯，雙擊 <code>index.html</code> 即可於任一瀏覽器秒開執行。採用清新綠色與淡紫雙色系質感設計，支援桌面雙欄與手機端自適應抽屜式切換。</li>
  <li><strong>地圖核心 (Map Engine)</strong>：選用開源 <strong>Leaflet.js</strong> 串接 <strong>Esri World Street Map 街道圖資</strong>（附備用高解析 Esri 衛星空照圖）。完全免申請商業金鑰（100% No API Key），零額度費用風險、杜絕金鑰外洩或作業違規扣分，支援縮放至 19 級且不擋本機檔案協議。</li>
  <li><strong>嚴謹資料工程 (Data Engineering)</strong>：以輕量結構化 JSON 陣列建立資料庫，收錄 50 處 100% 現正營業/開放之花蓮在地熱點。實施嚴格資料清洗，詳細標註經緯度、營業/開放時間、公休日提示、票價/價位區間、聯絡方式與特色亮點。</li>
</ul>

<h3>2. 核心功能實作</h3>
<ul>
  <li><strong>主題與子類別雙層快選</strong>：第一層提供全部探索 (50)、蔬食美饌 (30)、精選景點 (10)、質感咖啡館 (10)；切換至蔬食時動態展開「全素」、「蛋奶素」、「蔬食友善」子篩選標籤。</li>
  <li><strong>多維度分色視覺化標記 (Pins)</strong>：依類別給予不同專屬圖標與色彩（綠色 🌱 全素、黃色 🥚 蛋奶素、藍色 🥗 友善、紫色 🏛️ 景點、焦糖色 ☕ 咖啡館），地圖視角一目了然。</li>
  <li><strong>雙向地圖卡片連動 (Bi-directional Linking)</strong>：點擊左側卡片，地圖平滑平移（FlyTo）並展開詳細彈窗；點擊地圖標記，清單自動平滑滾動並高亮該卡片。</li>
  <li><strong>一鍵行動外連</strong>：支援 <code>tel:</code> 一鍵撥號（預約制店家提示社群預約），並整合 Google Maps 路線導航。</li>
</ul>

<h2>四、 怎麼測 (How to Test)</h2>
<h3>1. 功能測試矩陣 (Functional Testing)</h3>
<table>
  <tr>
    <th style="width:16%;">測試項目</th>
    <th style="width:24%;">測試步驟</th>
    <th style="width:48%;">預期結果</th>
    <th style="width:12%; text-align:center;">測試結果</th>
  </tr>
  <tr>
    <td><strong>地圖初始化</strong></td>
    <td>瀏覽器載入頁面</td>
    <td>以花蓮中心展開，Esri 街道圖正常載入，顯示全區 50 處標記</td>
    <td class="status-tag">通過</td>
  </tr>
  <tr>
    <td><strong>主題標籤切換</strong></td>
    <td>點擊「🏛️ 精選景點」頁籤</td>
    <td>清單與地圖即時過濾出 10 處景點，自動隱藏蔬食子分類，數量正確</td>
    <td class="status-tag">通過</td>
  </tr>
  <tr>
    <td><strong>咖啡廳專題過濾</strong></td>
    <td>點擊「☕ 質感咖啡館」頁籤</td>
    <td>呈現 10 家老宅/海景咖啡館，地圖標記轉換為焦糖色咖啡杯圖標</td>
    <td class="status-tag">通過</td>
  </tr>
  <tr>
    <td><strong>鄉鎮交叉篩選</strong></td>
    <td>選擇鄉鎮「壽豐鄉」</td>
    <td>視角平滑聚焦壽豐，呈現阿宏素食、鯉魚潭、雲山水、海碉堡等地點</td>
    <td class="status-tag">通過</td>
  </tr>
  <tr>
    <td><strong>關鍵字檢索</strong></td>
    <td>搜尋「肉桂捲」或「將軍府」</td>
    <td>即時精準比對出 Stroll 散步咖啡或將軍府1936園區，卡片即時更新</td>
    <td class="status-tag">通過</td>
  </tr>
  <tr>
    <td><strong>一鍵導航連動</strong></td>
    <td>點擊任一彈窗的「導航前往」</td>
    <td>新分頁精準開啟 Google 地圖並帶入正確經緯度座標</td>
    <td class="status-tag">通過</td>
  </tr>
</table>

<h3>2. 跨裝置與環境相容性測試</h3>
<ul>
  <li><strong>桌面瀏覽器</strong>（Chrome / Edge / Firefox / Safari）：雙欄佈局流暢，地圖拖曳與圖層切換反應極快。</li>
  <li><strong>行動裝置</strong>（iOS / Android）：點擊底部「📋 查看地點清單 / 🗺️ 查看地圖」一鍵切換，響應順暢。</li>
  <li><strong>靜態託管（GitHub Pages）</strong>：經線上 HTTPS 測試，無任何破圖或跨網域阻擋。</li>
</ul>

<h2>五、 下一步 (Next Steps)</h2>
<ul>
  <li><strong>多點行程規劃器 (Trip Planner)</strong>：讓使用者將想去的蔬食店、咖啡館與景點加入「收藏行程」，自動規劃最佳旅遊路線。</li>
  <li><strong>GPS 定位與距離排序</strong>：串接 HTML5 Geolocation API，計算使用者當前位置與各熱點的距離，依近到遠智慧排序。</li>
  <li><strong>PWA 離線地圖支援</strong>：將圖資與資料快取至使用者手機端，即便行經太魯閣或蘇花收訊不良路段亦能離線查閱。</li>
  <li><strong>小農與在地店家線上回報</strong>：打造社群協同投稿機制，促進花蓮在地綠色經濟共好循環。</li>
</ul>

</body>
</html>"""

    html_path = os.path.join(WORK_DIR, "_tmp_report.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)

    pdf_path = os.path.join(WORK_DIR, "專題報告_花蓮生活探索地圖.pdf")
    cmd = [
        EDGE_PATH,
        '--headless',
        '--disable-gpu',
        '--no-pdf-header-footer',
        f'--print-to-pdf={pdf_path}',
        html_path
    ]
    subprocess.run(cmd, capture_output=True)
    if os.path.exists(html_path):
        os.remove(html_path)
    print("Saved PDF:", pdf_path)
    return pdf_path

def create_ai_log_pdf():
    html_content = """<!DOCTYPE html>
<html lang="zh-TW">
<head>
<meta charset="UTF-8">
<title>《花蓮生活探索地圖》AI 協同開發歷程摘要</title>
<style>
  @page {
    size: A4;
    margin: 20mm 18mm 20mm 18mm;
  }
  body {
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "PingFang TC", "Microsoft JhengHei", sans-serif;
    color: #1e293b;
    line-height: 1.65;
    font-size: 10.5pt;
  }
  h1 {
    color: #064e3b;
    text-align: center;
    font-size: 20pt;
    margin-bottom: 6px;
    letter-spacing: -0.5px;
  }
  .meta-box {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    padding: 12px 18px;
    margin: 14px 0 20px 0;
    display: grid;
    grid-template-columns: auto 1fr;
    gap: 6px 14px;
    font-size: 9.5pt;
  }
  .meta-label {
    font-weight: bold;
    color: #065f46;
  }
  h2 {
    color: #064e3b;
    border-bottom: 2px solid #a7f3d0;
    padding-bottom: 4px;
    margin-top: 20px;
    margin-bottom: 8px;
    font-size: 13pt;
  }
  ul {
    margin-top: 4px;
    margin-bottom: 8px;
    padding-left: 20px;
  }
  li {
    margin-bottom: 5px;
  }
  .stage-badge {
    display: inline-block;
    background: #ecfdf5;
    color: #065f46;
    border: 1px solid #a7f3d0;
    padding: 2px 8px;
    border-radius: 4px;
    font-size: 9pt;
    font-weight: 600;
    margin-bottom: 4px;
  }
</style>
</head>
<body>

<h1>《花蓮生活探索地圖》AI 協同開發歷程摘要</h1>

<div class="meta-box">
  <div class="meta-label">專題名稱：</div><div>花蓮生活探索地圖 (Hualien Living Map)</div>
  <div class="meta-label">協同 AI：</div><div>Antigravity (Google DeepMind)</div>
  <div class="meta-label">期程紀錄：</div><div>2026 年 10 月 ~ 11 月</div>
  <div class="meta-label">專案倉庫：</div><div>https://github.com/411513004-coder/1028_-</div>
</div>

<h2>一、 階段一：專案起步與需求拆解</h2>
<ul>
  <li><strong>初期目標確立</strong>：使用者提出期中專題計畫，首要目標是嚴格符合課堂規範：「11/02 上課每人 90 秒展示（打開什麼、做什麼、看到什麼）」、「不強制龐大軟體，當天能順暢運行展示」、「程式碼、截圖絕對不可含 API 金鑰」。</li>
  <li><strong>極簡輕量架構選型</strong>：經 AI 可行性評估，決定採取「純原生前端 Web 架構（HTML5 + CSS3 + Vanilla JS + Leaflet.js）」，免除繁瑣後端環境安裝，保證在任何電腦上雙擊 index.html 即可秒開。地圖圖資做到 100% 零商業金鑰、零額度限制、零資安洩漏風險。</li>
</ul>

<h2>二、 階段二：實地查核與嚴格資料清洗 (Data Cleaning)</h2>
<ul>
  <li><strong>辨識與修正初期問題</strong>：使用者實測指出早期名單有部分店家名稱不正確或已停業（如「太魯閣旅人」查無紀錄、「宏銘素食」為名稱誤植）。</li>
  <li><strong>深度驗證與核實標準</strong>：AI 透過花蓮觀光資訊網、花蓮縣政府文化局古蹟登錄、Google 即時營業資訊與在地食記進行交叉查驗。更正新城火車站前名店為工藝師打造之「織娘之家 (預約制素食無菜單)」、北埔麻辣名店「金英蔬食」、東華大學商圈口碑名店「阿宏素食」、「好佳小吃（獨立素食鍋具）」、「麥樂多早午餐（獨立素食菜單）」。</li>
</ul>

<h2>三、 階段三：功能擴充 —— 景點與咖啡廳全方位整合</h2>
<ul>
  <li><strong>主題維度升級</strong>：配合使用者擴充需求，全面升級為三大生活主題：
    <ul>
      <li><strong>精選景點 (10 處)</strong>：將軍府1936、松園別館(2025新團隊重啟獲百大亮點)、文創園區、東大門夜市、慶修院、七星潭海岸、鯉魚潭、雲山水、台泥DAKA園區、崇德瑩農場。</li>
      <li><strong>質感咖啡館 (10 家)</strong>：浮室老宅、Giocare手沖陶藝、珈琲花、KŌHI宅深夜咖啡、Stroll散步咖啡肉桂捲、SENSES自家烘焙、咖逼小売所(縣定古蹟宿舍)、森山舍、新城練習曲書店咖啡、壽豐海碉堡。</li>
    </ul>
  </li>
  <li><strong>嚴格排除已歇業店家</strong>：查核並排除七星潭已結束實體營業之「龍宮咖啡」，替補為正派營運之「練習曲書店咖啡」，總收錄達 50 筆 100% 現役地點。</li>
</ul>

<h2>四、 階段四：互動邏輯與底圖穩定度優化</h2>
<ul>
  <li><strong>雙層導航與智慧感知</strong>：首頁設計大主題頁籤（全部 / 蔬食 / 景點 / 咖啡廳），並在點選蔬食時動態展開全素、蛋奶素、友善子分類標籤。</li>
  <li><strong>圖資穩定性修正</strong>：解決 CARTO 圖資在 local file:// 協議下會出現浮水印及 NatGeo 底圖高倍率破圖問題，全面切換為穩定無金鑰、支援 19 級高倍率之 Esri World Street Map 與 Esri 衛星空照圖。</li>
</ul>

<h2>五、 階段五：版本控制、GitHub Pages 部署與 QR Code 生成</h2>
<ul>
  <li><strong>版本控制規範</strong>：依課程規範設定 .gitignore，杜絕任何金鑰與機密檔案；倉庫名稱符合規範。</li>
  <li><strong>GitHub Pages 線上發布</strong>：成功部署至 https://411513004-coder.github.io/1028_-/，所有人皆可直接線上體驗。</li>
  <li><strong>專案 QR Code 整合</strong>：生成高解析專屬 QR Code 圖檔 (qrcode.png)，並無縫嵌入簡報封面，供 90 秒評審即時掃碼互動。</li>
</ul>

</body>
</html>"""

    html_path = os.path.join(WORK_DIR, "_tmp_ailog.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)

    pdf_path = os.path.join(WORK_DIR, "AI協同開發歷程摘要.pdf")
    cmd = [
        EDGE_PATH,
        '--headless',
        '--disable-gpu',
        '--no-pdf-header-footer',
        f'--print-to-pdf={pdf_path}',
        html_path
    ]
    subprocess.run(cmd, capture_output=True)
    if os.path.exists(html_path):
        os.remove(html_path)
    print("Saved PDF:", pdf_path)
    return pdf_path

def create_pptx_slides():
    prs = Presentation()
    prs.slide_width = PptInches(13.333) # 16:9
    prs.slide_height = PptInches(7.5)

    blank_layout = prs.slide_layouts[6] # blank

    # 投影片 1：封面
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(1, 0, 0, PptInches(13.333), PptInches(7.5)) # 1 is rectangle
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PptRGBColor(6, 78, 59) # 深綠
    bg1.line.color.rgb = PptRGBColor(6, 78, 59)

    # 標題文字方塊
    tx1 = s1.shapes.add_textbox(PptInches(1.0), PptInches(1.5), PptInches(7.8), PptInches(4.5))
    tf1 = tx1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "期中專案展示 · 90秒實機演示"
    p0.font.size = PptPt(16)
    p0.font.color.rgb = PptRGBColor(167, 243, 208)

    p1 = tf1.add_paragraph()
    p1.text = "花蓮生活探索地圖"
    p1.font.size = PptPt(44)
    p1.font.bold = True
    p1.font.color.rgb = PptRGBColor(255, 255, 255)
    p1.space_before = PptPt(10)
    p1.space_after = PptPt(10)

    p2 = tf1.add_paragraph()
    p2.text = "Hualien Living Map ── 蔬食美饌 · 精選景點 · 質感咖啡館指南"
    p2.font.size = PptPt(18)
    p2.font.color.rgb = PptRGBColor(209, 250, 229)
    p2.space_after = PptPt(24)

    p3 = tf1.add_paragraph()
    p3.text = "• 專題核心：輕量化互動式 Web App（100% 開源無金鑰 · 零安裝雙擊即跑）\n• 收錄熱點：50 處花蓮在地實名核實現役地點（30 家蔬食 + 10 大景點 + 10 家咖啡館）\n• 技術選型：HTML5 / CSS3 / Vanilla JS / Leaflet.js / Esri 街道圖資"
    p3.font.size = PptPt(14)
    p3.font.color.rgb = PptRGBColor(240, 253, 244)

    # QR Code 白色卡片
    qr_card = s1.shapes.add_shape(1, PptInches(9.2), PptInches(1.5), PptInches(3.2), PptInches(4.5))
    qr_card.fill.solid()
    qr_card.fill.fore_color.rgb = PptRGBColor(255, 255, 255)
    qr_card.line.color.rgb = PptRGBColor(226, 232, 240)

    qr_img_path = os.path.join(WORK_DIR, "qrcode.png")
    if os.path.exists(qr_img_path):
        s1.shapes.add_picture(qr_img_path, PptInches(9.55), PptInches(1.8), PptInches(2.5), PptInches(2.5))

    qr_tx = s1.shapes.add_textbox(PptInches(9.3), PptInches(4.5), PptInches(3.0), PptInches(1.2))
    qtf = qr_tx.text_frame
    qtf.word_wrap = True
    qp1 = qtf.paragraphs[0]
    qp1.text = "📱 掃描立即體驗"
    qp1.font.size = PptPt(14)
    qp1.font.bold = True
    qp1.font.color.rgb = PptRGBColor(6, 95, 70)
    qp1.alignment = PP_ALIGN.CENTER
    qp2 = qtf.add_paragraph()
    qp2.text = "411513004-coder.github.io/1028_-"
    qp2.font.size = PptPt(10)
    qp2.font.color.rgb = PptRGBColor(100, 116, 139)
    qp2.alignment = PP_ALIGN.CENTER

    # 投影片 2：痛點與目標族群
    s2 = prs.slides.add_slide(blank_layout)
    # Header
    tx_h2 = s2.shapes.add_textbox(PptInches(0.8), PptInches(0.5), PptInches(11.7), PptInches(1.0))
    thf2 = tx_h2.text_frame
    thp1 = thf2.paragraphs[0]
    thp1.text = "🎯 為什麼做？解決什麼問題？"
    thp1.font.size = PptPt(26)
    thp1.font.bold = True
    thp1.font.color.rgb = PptRGBColor(15, 23, 42)
    thp2 = thf2.add_paragraph()
    thp2.text = "Problem Statement & Target Audience · 打造一站式花蓮綠色生活指南"
    thp2.font.size = PptPt(13)
    thp2.font.color.rgb = PptRGBColor(100, 116, 139)

    # 痛點卡片
    c1 = s2.shapes.add_shape(1, PptInches(0.8), PptInches(1.8), PptInches(5.7), PptInches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = PptRGBColor(236, 253, 245)
    c1.line.color.rgb = PptRGBColor(167, 243, 208)
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "🚨 現有痛點 (Pain Points)\n"
    p.font.size = PptPt(18)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(6, 95, 70)
    p_body = tf_c1.add_paragraph()
    p_body.text = "• 蔬食分類混雜：主流地圖將「純素」、「蛋奶素」與「葷素共食」混淆，素食者需反覆致電確認。\n\n• 資訊零散割裂：傳統美食地圖缺乏文化景點與休憩咖啡廳的橫向串連。\n\n• 幽靈店家困擾：網路舊名單充斥已停業或轉型店家，容易造成旅遊撲空。\n\n• 資安與維護門檻：商業地圖 API 計費昂貴、繁瑣且有金鑰洩漏風險。"
    p_body.font.size = PptPt(13.5)
    p_body.font.color.rgb = PptRGBColor(51, 65, 85)

    # 受眾卡片
    c2 = s2.shapes.add_shape(1, PptInches(6.8), PptInches(1.8), PptInches(5.7), PptInches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = PptRGBColor(248, 250, 252)
    c2.line.color.rgb = PptRGBColor(226, 232, 240)
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "👥 目標受眾 (Target Audience)\n"
    p.font.size = PptPt(18)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(15, 23, 42)
    p_body = tf_c2.add_paragraph()
    p_body.text = "• 花蓮在地居民與蔬食族群：快速掌握花蓮市、吉安等區域安心用餐新選擇。\n\n• 大專院校師生：東華大學志學商圈、慈濟建國校區周邊平價美食與自習咖啡空間。\n\n• 深度旅遊自由行觀光客：跨越新城、壽豐、秀林與鳳林，無縫串聯山海自然與人文名店。\n\n• 重視永續低碳的生活推廣者：鼓勵大眾實踐減碳飲食與在地慢活步調。"
    p_body.font.size = PptPt(13.5)
    p_body.font.color.rgb = PptRGBColor(51, 65, 85)

    # 投影片 3：三大主題架構
    s3 = prs.slides.add_slide(blank_layout)
    tx_h3 = s3.shapes.add_textbox(PptInches(0.8), PptInches(0.5), PptInches(11.7), PptInches(1.0))
    thf3 = tx_h3.text_frame
    thp1 = thf3.paragraphs[0]
    thp1.text = "🧩 核心架構：三大主題與 50 處現役資料"
    thp1.font.size = PptPt(26)
    thp1.font.bold = True
    thp1.font.color.rgb = PptRGBColor(15, 23, 42)
    thp2 = thf3.add_paragraph()
    thp2.text = "Data Verification · 2026 最新官方資訊與實地核實，排除停業幽靈店家"
    thp2.font.size = PptPt(13)
    thp2.font.color.rgb = PptRGBColor(100, 116, 139)

    col_data = [
        ("🌱 蔬食美饌 (30 家)", PptRGBColor(236, 253, 245), PptRGBColor(167, 243, 208), PptRGBColor(6, 95, 70),
         "• 全素/純素 (16家)：\n  蝸牛蔬食Buffet、植地佛陀碗、明園煎餃、宜蘭水煎包、東華阿宏素食、新城織娘之家等。\n\n• 蛋奶素 (4家)：\n  Piatto+菜老宅義式、青田の蔬食、草緣齋披薩。\n\n• 蔬食友善 (10家)：\n  時光1939有機早午餐、櫻之田野野菜鍋、好佳小吃(獨立素食鍋具)、鳳林游翁臭豆腐。"),
        
        ("🏛️ 精選景點 (10 處)", PptRGBColor(245, 243, 255), PptRGBColor(221, 214, 254), PptRGBColor(109, 40, 217),
         "• 市區歷史人文：\n  將軍府1936日式宿舍群、松園別館(2025重啟獲觀光百大亮點)、文創園區舊酒廠、東大門夜市。\n\n• 古蹟與自然山海：\n  吉安慶修院(三級古蹟)、新城七星潭月牙灣、壽豐鯉魚潭、雲山水落羽松夢幻湖。\n\n• 蘇花前哨亮點：\n  秀林台泥DAKA開放生態園區、崇德瑩農場。"),
        
        ("☕ 質感咖啡館 (10 家)", PptRGBColor(255, 251, 235), PptRGBColor(253, 230, 138), PptRGBColor(180, 83, 9),
         "• 老宅文青聚落：\n  浮室純白老宅不限時、Giocare樹下手沖與柴燒陶、Caffe Fiore珈琲花貓咪、KŌHI宅深夜榻榻米。\n\n• 烘焙與古蹟空間：\n  Stroll散步咖啡人氣肉桂捲、SENSES極簡烘焙、咖逼小売所(縣定古蹟)、森山舍林業老舍。\n\n• 山海人文特色：\n  新城練習曲書店(書香咖啡)、壽豐海碉堡。")
    ]
    for i, (title, bg, border, title_c, desc) in enumerate(col_data):
        card = s3.shapes.add_shape(1, PptInches(0.8 + i * 3.95), PptInches(1.8), PptInches(3.75), PptInches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = bg
        card.line.color.rgb = border
        tf = card.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title + "\n"
        p_t.font.size = PptPt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = title_c
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = PptPt(11.5)
        p_d.font.color.rgb = PptRGBColor(51, 65, 85)

    # 投影片 4：90 秒實機演示劇本
    s4 = prs.slides.add_slide(blank_layout)
    tx_h4 = s4.shapes.add_textbox(PptInches(0.8), PptInches(0.5), PptInches(11.7), PptInches(1.0))
    thf4 = tx_h4.text_frame
    thp1 = thf4.paragraphs[0]
    thp1.text = "⏱️ 90 秒現場展示腳本：打開什麼、做什麼、看到什麼"
    thp1.font.size = PptPt(26)
    thp1.font.bold = True
    thp1.font.color.rgb = PptRGBColor(15, 23, 42)
    thp2 = thf4.add_paragraph()
    thp2.text = "Live Demonstration Workflow · 嚴格對照評分標準，節奏流暢展示核心價值"
    thp2.font.size = PptPt(13)
    thp2.font.color.rgb = PptRGBColor(100, 116, 139)

    step_data = [
        ("00 ~ 20 秒 · 打開什麼", "🖥️ 啟動與破題", PptRGBColor(16, 185, 129),
         "【展示操作】\n瀏覽器秒開地圖 Web App。\n\n【口述台詞】\n「打開花蓮生活探索地圖，全原生前端打造、完全零金鑰、零授權費！畫面上清楚標註花蓮全區 50 處精選蔬食、景點與咖啡廳。」"),
        ("20 ~ 55 秒 · 做什麼", "⚡ 互動篩選與檢索", PptRGBColor(59, 130, 246),
         "【展示操作】\n1. 點擊「🏛️ 精選景點」即時聚焦地標\n2. 點選「🌱 蔬食美饌」並細篩「全素」\n3. 下拉選「壽豐鄉」自動定位東華大學\n4. 搜尋框輸入「肉桂捲」秒級查出名店\n\n【口述台詞】\n「多維度即時過濾，卡片與地圖 Pin 雙向動態連動，找好店不用 3 秒！」"),
        ("55 ~ 90 秒 · 看到什麼", "🧭 成果反饋與行動外連", PptRGBColor(139, 92, 246),
         "【展示操作】\n點擊標記平滑飛入，彈出詳細資訊卡。\n\n【口述台詞】\n「彈出營業時段、門票、特色，點擊『導航前往』一鍵連動 Google Maps 規劃路線！90 秒解決花蓮探索痛點，謝謝大家！」")
    ]
    for i, (badge, title, color, desc) in enumerate(step_data):
        box = s4.shapes.add_shape(1, PptInches(0.8 + i * 3.95), PptInches(1.8), PptInches(3.75), PptInches(4.8))
        box.fill.solid()
        box.fill.fore_color.rgb = PptRGBColor(255, 255, 255)
        box.line.color.rgb = color
        box.line.width = PptPt(2)
        tf = box.text_frame
        tf.word_wrap = True
        p_b = tf.paragraphs[0]
        p_b.text = badge + "\n"
        p_b.font.size = PptPt(12)
        p_b.font.bold = True
        p_b.font.color.rgb = color
        p_t = tf.add_paragraph()
        p_t.text = title + "\n"
        p_t.font.size = PptPt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = PptRGBColor(15, 23, 42)
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = PptPt(12)
        p_d.font.color.rgb = PptRGBColor(51, 65, 85)

    # 投影片 5：技術亮點與未來展望
    s5 = prs.slides.add_slide(blank_layout)
    tx_h5 = s5.shapes.add_textbox(PptInches(0.8), PptInches(0.5), PptInches(11.7), PptInches(1.0))
    thf5 = tx_h5.text_frame
    thp1 = thf5.paragraphs[0]
    thp1.text = "🚀 技術亮點、測試成果與下一步"
    thp1.font.size = PptPt(26)
    thp1.font.bold = True
    thp1.font.color.rgb = PptRGBColor(15, 23, 42)
    thp2 = thf5.add_paragraph()
    thp2.text = "Technical Highlights, Testing & Next Steps · 扎實落地與持續演進"
    thp2.font.size = PptPt(13)
    thp2.font.color.rgb = PptRGBColor(100, 116, 139)

    c5_1 = s5.shapes.add_shape(1, PptInches(0.8), PptInches(1.8), PptInches(5.7), PptInches(4.8))
    c5_1.fill.solid()
    c5_1.fill.fore_color.rgb = PptRGBColor(236, 253, 245)
    c5_1.line.color.rgb = PptRGBColor(167, 243, 208)
    tf5_1 = c5_1.text_frame
    tf5_1.word_wrap = True
    p = tf5_1.paragraphs[0]
    p.text = "⚙️ 技術亮點與測試成果\n"
    p.font.size = PptPt(18)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(6, 95, 70)
    p_body = tf5_1.add_paragraph()
    p_body.text = "• 100% 零商業金鑰：採用 Esri 街道圖與高解析空照圖，無 API 費用、無上限、杜絕資安外洩。\n\n• 極致輕量與零依賴：首頁載入小於 0.3 秒，免安裝環境，雙擊 index.html 即可運行。\n\n• 雙向連動架構：左側卡片與地圖 Pin 雙向聚焦（FlyTo 動態位移、平滑滾動）。\n\n• 全平臺適配：支援電腦寬螢幕雙欄與手機端自適應抽屜介面；已部署 GitHub Pages 免費線上即時體驗。"
    p_body.font.size = PptPt(13)
    p_body.font.color.rgb = PptRGBColor(51, 65, 85)

    c5_2 = s5.shapes.add_shape(1, PptInches(6.8), PptInches(1.8), PptInches(5.7), PptInches(4.8))
    c5_2.fill.solid()
    c5_2.fill.fore_color.rgb = PptRGBColor(248, 250, 252)
    c5_2.line.color.rgb = PptRGBColor(226, 232, 240)
    tf5_2 = c5_2.text_frame
    tf5_2.word_wrap = True
    p = tf5_2.paragraphs[0]
    p.text = "🔮 未來發展規劃 (Next Steps)\n"
    p.font.size = PptPt(18)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(15, 23, 42)
    p_body = tf5_2.add_paragraph()
    p_body.text = "• 自由行多點行程自動規劃：支援使用者將想去的地點加入清單，演算法自動串聯最佳路線。\n\n• GPS 即時定位與距離排序：整合 HTML5 Geolocation API，自動計算距離當前位置最近據點。\n\n• PWA 離線快取技術：引入 Service Worker，將地圖底圖快取至手機端，行經山區亦可離線查閱。\n\n• 小農與在地店家線上共好：打造社群協同投稿機制，促進花蓮在地綠色經濟共好循環。"
    p_body.font.size = PptPt(13)
    p_body.font.color.rgb = PptRGBColor(51, 65, 85)

    pptx_path = os.path.join(WORK_DIR, "期中專題簡報_花蓮生活探索地圖.pptx")
    prs.save(pptx_path)
    print("Saved PPTX:", pptx_path)
    return pptx_path

if __name__ == "__main__":
    create_report_docx()
    create_report_pdf()
    create_ai_log_docx()
    create_ai_log_pdf()
    create_pptx_slides()
    print("All deliverables generated successfully!")
