import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor
from pptx.enum.text import PP_ALIGN

WORK_DIR = r"c:\User\Downloads\1023ai期中"
EDGE_PATH = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
QR_IMG_PATH = os.path.join(WORK_DIR, "qrcode.png")

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

# ==========================================
# 1. 建立 4:3 PowerPoint (.pptx)
# ==========================================
def create_43_pptx():
    prs = Presentation()
    prs.slide_width = PptInches(10.0)   # 4:3 寬度
    prs.slide_height = PptInches(7.5)  # 4:3 高度

    blank_layout = prs.slide_layouts[6]

    # Helper: 加頁面頂部 Header
    def add_slide_header(slide, title, subtitle):
        tb = slide.shapes.add_textbox(PptInches(0.6), PptInches(0.4), PptInches(8.8), PptInches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = PptPt(22)
        p1.font.bold = True
        p1.font.color.rgb = PptRGBColor(15, 23, 42)
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = PptPt(11)
        p2.font.color.rgb = PptRGBColor(100, 116, 139)

    # ------------------------------------------
    # 頁面 1：封面 (4:3)
    # ------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(1, 0, 0, PptInches(10.0), PptInches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PptRGBColor(6, 78, 59)
    bg1.line.color.rgb = PptRGBColor(6, 78, 59)

    tb1 = s1.shapes.add_textbox(PptInches(0.8), PptInches(1.2), PptInches(5.6), PptInches(5.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "期中專題 · 開發歷程與人機協作全紀錄"
    p0.font.size = PptPt(13)
    p0.font.bold = True
    p0.font.color.rgb = PptRGBColor(167, 243, 208)

    p1 = tf1.add_paragraph()
    p1.text = "花蓮生活探索地圖"
    p1.font.size = PptPt(36)
    p1.font.bold = True
    p1.font.color.rgb = PptRGBColor(255, 255, 255)
    p1.space_before = PptPt(8)
    p1.space_after = PptPt(8)

    p2 = tf1.add_paragraph()
    p2.text = "從零發想、踩坑排雷到最終落地實踐"
    p2.font.size = PptPt(16)
    p2.font.color.rgb = PptRGBColor(209, 250, 229)
    p2.space_after = PptPt(18)

    p3 = tf1.add_paragraph()
    p3.text = "• 專題主題：蔬食美饌 · 精選景點 · 質感咖啡館指南\n• 技術核心：HTML5 / Vanilla JS / Leaflet / Esri 圖資\n• 協同開發：Antigravity (Google DeepMind) 協作\n• 部署網址：411513004-coder.github.io/1028_-"
    p3.font.size = PptPt(11.5)
    p3.font.color.rgb = PptRGBColor(240, 253, 244)

    # 封面右側 QR Code 卡片
    qr_card = s1.shapes.add_shape(1, PptInches(6.8), PptInches(1.5), PptInches(2.6), PptInches(4.2))
    qr_card.fill.solid()
    qr_card.fill.fore_color.rgb = PptRGBColor(255, 255, 255)
    qr_card.line.color.rgb = PptRGBColor(226, 232, 240)

    if os.path.exists(QR_IMG_PATH):
        s1.shapes.add_picture(QR_IMG_PATH, PptInches(7.1), PptInches(1.8), PptInches(2.0), PptInches(2.0))

    q_tb = s1.shapes.add_textbox(PptInches(6.9), PptInches(4.0), PptInches(2.4), PptInches(1.2))
    q_tf = q_tb.text_frame
    q_tf.word_wrap = True
    qp1 = q_tf.paragraphs[0]
    qp1.text = "📱 掃碼即刻體驗"
    qp1.font.size = PptPt(12)
    qp1.font.bold = True
    qp1.font.color.rgb = PptRGBColor(6, 95, 70)
    qp1.alignment = PP_ALIGN.CENTER
    qp2 = q_tf.add_paragraph()
    qp2.text = "411513004-coder.github.io/1028_-"
    qp2.font.size = PptPt(8.5)
    qp2.font.color.rgb = PptRGBColor(100, 116, 139)
    qp2.alignment = PP_ALIGN.CENTER

    # ------------------------------------------
    # 頁面 2：專案起點與初始構想
    # ------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_header(s2, "🚀 專案起點：為什麼做？如何出發？", "Project Inception & Core Requirements · 兼顧課堂規範與真實生活需求")

    c1 = s2.shapes.add_shape(1, PptInches(0.6), PptInches(1.5), PptInches(4.2), PptInches(5.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = PptRGBColor(236, 253, 245)
    c1.line.color.rgb = PptRGBColor(167, 243, 208)
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "📌 課堂要求與現實限制\n"
    p.font.size = PptPt(15)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(6, 95, 70)
    p_body = tf1.add_paragraph()
    p_body.text = "• 90 秒即開即評：課堂要求展示打開什麼、做什麼、看到什麼，完全不容許展示時環境崩潰或等待載入。\n\n• 零金鑰資安紅線：程式碼、截圖與遠端倉庫嚴禁含有任何 API 金鑰，否則重扣。\n\n• 零安裝零負擔：避免笨重伺服器架構，任一電腦瀏覽器雙擊即可順暢運作。"
    p_body.font.size = PptPt(11)
    p_body.font.color.rgb = PptRGBColor(51, 65, 85)

    c2 = s2.shapes.add_shape(1, PptInches(5.2), PptInches(1.5), PptInches(4.2), PptInches(5.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = PptRGBColor(248, 250, 252)
    c2.line.color.rgb = PptRGBColor(226, 232, 240)
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "💡 初衷發想：花蓮綠色生活指南\n"
    p.font.size = PptPt(15)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(15, 23, 42)
    p_body = tf2.add_paragraph()
    p_body.text = "• 解決素食痛點：主流地圖把純素、蛋奶素與葷素共食混為一談，常讓素食者白跑或反覆打電話。\n\n• 學生與旅客剛需：東華大學、慈濟生活圈缺乏整合性在地地圖，外地旅客常找不到靠譜推薦。\n\n• 決定尋求 AI 協同：由 Antigravity 快速產出系統雛形與前端骨架，開啟人機協作之旅。"
    p_body.font.size = PptPt(11)
    p_body.font.color.rgb = PptRGBColor(51, 65, 85)

    # ------------------------------------------
    # 頁面 3：遭遇困難一：AI 幻覺與幽靈店家
    # ------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_header(s3, "🚨 遭遇困難一：AI 幻覺、幽靈店家與店名自創", "Challenge 1: AI Hallucinations & Ghost Stores · 速度飛快，但資料暗藏危機")

    h_data = [
        ("👻 幽靈店家：抓取已停業名單", PptRGBColor(254, 242, 242), PptRGBColor(254, 202, 202), PptRGBColor(185, 28, 28),
         "【現象】AI 剛開始蒐集資料速度極快，但在細查後發現：\n• 七星潭人氣老店「龍宮咖啡」官方 IG 早已宣布歇業！\n• 搜尋出現早已不存在的歷史店家。\n\n【排雷】人工覆核最新社群動態，剔除停業店家，替補新城老街正派經營之「練習曲書店咖啡」。"),
        ("🏷️ 憑空捏造：幫現存店家創生新名字", PptRGBColor(255, 247, 237), PptRGBColor(254, 215, 170), PptRGBColor(194, 65, 12),
         "【現象】店家明明存在，AI 卻自己給了「平行時空新名稱」：\n• 查無任何登記紀錄的「太魯閣旅人素食小館」。\n• 壽豐東華商圈將「阿宏素食」誤植為「宏銘素食」。\n\n【排雷】透過在地生活經驗比對花蓮縣政府登錄資料，將新城名店更正為工藝師名店「織娘之家無菜單料理」。")
    ]
    for i, (title, bg, border, tc, desc) in enumerate(h_data):
        box = s3.shapes.add_shape(1, PptInches(0.6 + i * 4.5), PptInches(1.5), PptInches(4.3), PptInches(5.2))
        box.fill.solid()
        box.fill.fore_color.rgb = bg
        box.line.color.rgb = border
        tf = box.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title + "\n"
        p_t.font.size = PptPt(13.5)
        p_t.font.bold = True
        p_t.font.color.rgb = tc
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = PptPt(10.5)
        p_d.font.color.rgb = PptRGBColor(51, 65, 85)

    # ------------------------------------------
    # 頁面 4：遭遇困難二：地圖底圖踩坑與排錯
    # ------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_header(s4, "🗺️ 遭遇困難二：底圖破圖、浮水印與協議限制", "Challenge 2: Basemap Watermark & Zoom Breakdown · 零金鑰圖資的實戰考驗")

    m_data = [
        ("1. CARTO Voyager 浮水印限制", PptRGBColor(241, 245, 249), PptRGBColor(203, 213, 225), PptRGBColor(30, 41, 59),
         "【問題】在 local 端直接以 file:// 開啟 index.html 時，CARTO 圖資會自動覆蓋滿版「CARTO」水印，視覺受到嚴重干擾。\n\n【原因】CARTO 對無 Referrer 或本機檔案協議有防護限制。"),
        ("2. NatGeo 放大極限破圖", PptRGBColor(241, 245, 249), PptRGBColor(203, 213, 225), PptRGBColor(30, 41, 59),
         "【問題】嘗試換用 NatGeo 國家地理風圖資，但使用者拉近放大到市區街道（Zoom 16 以上）時，畫面全黑並提示「Map data not yet available」。\n\n【原因】該底圖最大僅支援至 16 級縮放。"),
        ("✅ 最終最佳解：Esri 雙圖層方案", PptRGBColor(236, 253, 245), PptRGBColor(167, 243, 208), PptRGBColor(6, 95, 70),
         "【解法】全面重構為 Esri 雙圖層架構：\n• 主底圖：Esri World Street Map（街道圖清晰無水印，支援高解析 Zoom 19 級）。\n• 副底圖：Esri World Imagery（高空衛星圖），100% 零商業金鑰、零付費、穩定度 100%。")
    ]
    for i, (title, bg, border, tc, desc) in enumerate(m_data):
        box = s4.shapes.add_shape(1, PptInches(0.6 + i * 2.95), PptInches(1.5), PptInches(2.8), PptInches(5.2))
        box.fill.solid()
        box.fill.fore_color.rgb = bg
        box.line.color.rgb = border
        tf = box.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title + "\n"
        p_t.font.size = PptPt(12.5)
        p_t.font.bold = True
        p_t.font.color.rgb = tc
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = PptPt(10)
        p_d.font.color.rgb = PptRGBColor(51, 65, 85)

    # ------------------------------------------
    # 頁面 5：功能擴充與三大主題成形
    # ------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_header(s5, "🌿 功能擴充：升級「花蓮生活探索地圖」", "Scope Expansion: From Veggie Map to Comprehensive Living Guide · 50 處全實名核實")

    cols5 = [
        ("🌱 蔬食美饌 (30 家)", PptRGBColor(236, 253, 245), PptRGBColor(167, 243, 208), PptRGBColor(6, 95, 70),
         "• 全素/純素 (16家)：\n  蝸牛蔬食Buffet、植地佛陀碗、明園煎餃、宜蘭水煎包、東華阿宏素食、織娘之家無菜單料理。\n• 蛋奶素 (4家)：\n  Piatto+菜老宅義式、青田の蔬食、草緣齋手工披薩。\n• 蔬食友善 (10家)：\n  時光1939有機早午餐、櫻之田野野菜鍋、好佳小吃(獨立素食鍋具)。"),
        ("🏛️ 精選景點 (10 處)", PptRGBColor(245, 243, 255), PptRGBColor(221, 214, 254), PptRGBColor(109, 40, 217),
         "• 人文歷史：\n  將軍府1936、松園別館(2025新團隊重啟獲百大亮點)、文創園區、東大門夜市、吉安慶修院。\n• 自然山海：\n  七星潭月牙灣、壽豐鯉魚潭、雲山水落羽松、秀林台泥DAKA園區、崇德瑩農場。"),
        ("☕ 質感咖啡館 (10 家)", PptRGBColor(255, 251, 235), PptRGBColor(253, 230, 138), PptRGBColor(180, 83, 9),
         "• 老宅特色：\n  浮室純白老宅、Giocare手沖陶藝、珈琲花乾燥花貓咪、KŌHI宅深夜榻榻米。\n• 烘焙與古蹟：\n  Stroll散步咖啡肉桂捲、SENSES烘焙、咖逼小売所(古蹟宿舍)、森山舍、練習曲書店、海碉堡。")
    ]
    for i, (title, bg, border, tc, desc) in enumerate(cols5):
        box = s5.shapes.add_shape(1, PptInches(0.6 + i * 2.95), PptInches(1.5), PptInches(2.8), PptInches(5.2))
        box.fill.solid()
        box.fill.fore_color.rgb = bg
        box.line.color.rgb = border
        tf = box.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title + "\n"
        p_t.font.size = PptPt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = tc
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = PptPt(9.8)
        p_d.font.color.rgb = PptRGBColor(51, 65, 85)

    # ------------------------------------------
    # 頁面 6：製作細節與技術亮點
    # ------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_header(s6, "⚙️ 製作細節：技術亮點與互動設計", "Technical Implementation & UI/UX Highlights · 打造流暢的使用者體驗")

    details = [
        ("雙層智慧導航系統", "首頁大標籤（全部 / 蔬食 / 景點 / 咖啡廳），並具備自動感知機制：僅在瀏覽蔬食時動態展開全素、蛋奶素、友善子分類標籤，乾淨不雜亂。"),
        ("雙向卡片地圖聯動", "點擊左側卡片，地圖平滑平移（FlyTo）並展開詳細彈窗；點擊地圖標記，清單自動平滑滾動並高亮該卡片，資訊直觀連貫。"),
        ("多維分色自訂標記", "綠色 🌱 全素、黃色 🥚 蛋奶素、藍色 🥗 友善、紫色 🏛️ 景點、焦糖色 ☕ 咖啡館，地圖全景一眼區隔，辨識效率極高。"),
        ("行動端無縫切換與外連", "支援手機自適應抽屜（一鍵切換清單與地圖）；整合 tel: 一鍵通話與 Google Maps GPS 經緯度路徑導航。")
    ]
    for i, (title, desc) in enumerate(details):
        r_idx = i // 2
        c_idx = i % 2
        card = s6.shapes.add_shape(1, PptInches(0.6 + c_idx * 4.5), PptInches(1.5 + r_idx * 2.65), PptInches(4.3), PptInches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = PptRGBColor(248, 250, 252)
        card.line.color.rgb = PptRGBColor(226, 232, 240)
        tf = card.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = f"✨ {title}\n"
        p_t.font.size = PptPt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = PptRGBColor(6, 95, 70)
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = PptPt(10.5)
        p_d.font.color.rgb = PptRGBColor(51, 65, 85)

    # ------------------------------------------
    # 頁面 7：倒數第二頁 · 我的心得（完整呈現使用者心得）
    # ------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_header(s7, "💬 專案心得與人機協作反思", "Reflection on AI-Assisted Development · 速度與真實性的深刻體悟")

    # 心得引言大卡片
    thought_card = s7.shapes.add_shape(1, PptInches(0.8), PptInches(1.5), PptInches(8.4), PptInches(3.2))
    thought_card.fill.solid()
    thought_card.fill.fore_color.rgb = PptRGBColor(236, 253, 245)
    thought_card.line.color.rgb = PptRGBColor(167, 243, 208)
    thought_card.line.width = PptPt(2)
    tf_t = thought_card.text_frame
    tf_t.word_wrap = True

    tp0 = tf_t.paragraphs[0]
    tp0.text = "「開發者的真實心聲」\n"
    tp0.font.size = PptPt(14)
    tp0.font.bold = True
    tp0.font.color.rgb = PptRGBColor(6, 95, 70)

    tp1 = tf_t.add_paragraph()
    tp1.text = "這次我請 Antigravity 幫我架設網站，架設的速度快，但是在修正資料正確性卻花了不少時間，他會找到已經停業的店，或者雖然店存在，但是幫它創造了新名字，還有地圖的問題等說不完的問題。\n\n在這次製作網站的過程中，雖然 AI 確實可以取代部份的工作，但是資料校正跟最後成品測試還是需要有人來把關。"
    tp1.font.size = PptPt(13.5)
    tp1.font.bold = True
    tp1.font.color.rgb = PptRGBColor(15, 23, 42)

    # 下方重點提煉 2 個小卡片
    sub1 = s7.shapes.add_shape(1, PptInches(0.8), PptInches(4.9), PptInches(4.05), PptInches(1.8))
    sub1.fill.solid()
    sub1.fill.fore_color.rgb = PptRGBColor(248, 250, 252)
    sub1.line.color.rgb = PptRGBColor(226, 232, 240)
    tf_s1 = sub1.text_frame
    tf_s1.word_wrap = True
    p = tf_s1.paragraphs[0]
    p.text = "⚡ AI 的優勢：架構與代碼產出神速"
    p.font.size = PptPt(12)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(16, 185, 129)
    p2 = tf_s1.add_paragraph()
    p2.text = "從無到有建立完整的 Leaflet 互動地圖、雙向連動、CSS 響應式佈局與 GitHub Pages 自動部署，省去數天寫代碼的摸索。"
    p2.font.size = PptPt(9.5)
    p2.font.color.rgb = PptRGBColor(71, 85, 105)

    sub2 = s7.shapes.add_shape(1, PptInches(5.15), PptInches(4.9), PptInches(4.05), PptInches(1.8))
    sub2.fill.solid()
    sub2.fill.fore_color.rgb = PptRGBColor(248, 250, 252)
    sub2.line.color.rgb = PptRGBColor(226, 232, 240)
    tf_s2 = sub2.text_frame
    tf_s2.word_wrap = True
    p = tf_s2.paragraphs[0]
    p.text = "🛡️ 人類的價值：嚴格把關與真實校驗"
    p.font.size = PptPt(12)
    p.font.bold = True
    p.font.color.rgb = PptRGBColor(234, 88, 12)
    p2 = tf_s2.add_paragraph()
    p2.text = "AI 缺乏真實世界驗證能力，容易把舊資料當現狀或自我造詞。唯有依靠人類的地圖實測、電話覆核與邏輯把關，才能產出真正有溫度的實用作品。"
    p2.font.size = PptPt(9.5)
    p2.font.color.rgb = PptRGBColor(71, 85, 105)

    # ------------------------------------------
    # 頁面 8：最後一頁 · 謝謝大家 (Thank You / Q&A)
    # ------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    bg8 = s8.shapes.add_shape(1, 0, 0, PptInches(10.0), PptInches(7.5))
    bg8.fill.solid()
    bg8.fill.fore_color.rgb = PptRGBColor(6, 78, 59)
    bg8.line.color.rgb = PptRGBColor(6, 78, 59)

    tb8 = s8.shapes.add_textbox(PptInches(1.0), PptInches(1.5), PptInches(5.0), PptInches(4.5))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    p0 = tf8.paragraphs[0]
    p0.text = "報告完畢 · 敬請指教"
    p0.font.size = PptPt(16)
    p0.font.bold = True
    p0.font.color.rgb = PptRGBColor(167, 243, 208)

    p1 = tf8.add_paragraph()
    p1.text = "謝謝大家！"
    p1.font.size = PptPt(44)
    p1.font.bold = True
    p1.font.color.rgb = PptRGBColor(255, 255, 255)
    p1.space_before = PptPt(10)
    p1.space_after = PptPt(14)

    p2 = tf8.add_paragraph()
    p2.text = "歡迎掃描右側 QR Code 現場體驗，或在 GitHub 查看完整無金鑰純原生開原始碼。\n\n🔗 專案首頁：\nhttps://411513004-coder.github.io/1028_-/\n\n📦 開源倉庫：\nhttps://github.com/411513004-coder/1028_-"
    p2.font.size = PptPt(12)
    p2.font.color.rgb = PptRGBColor(209, 250, 229)

    # 謝謝大家頁面的 QR Code
    qr_card8 = s8.shapes.add_shape(1, PptInches(6.4), PptInches(1.6), PptInches(2.8), PptInches(4.2))
    qr_card8.fill.solid()
    qr_card8.fill.fore_color.rgb = PptRGBColor(255, 255, 255)
    qr_card8.line.color.rgb = PptRGBColor(226, 232, 240)

    if os.path.exists(QR_IMG_PATH):
        s8.shapes.add_picture(QR_IMG_PATH, PptInches(6.7), PptInches(1.9), PptInches(2.2), PptInches(2.2))

    q_tb8 = s8.shapes.add_textbox(PptInches(6.5), PptInches(4.2), PptInches(2.6), PptInches(1.2))
    q_tf8 = q_tb8.text_frame
    q_tf8.word_wrap = True
    qp1 = q_tf8.paragraphs[0]
    qp1.text = "花蓮生活探索地圖"
    qp1.font.size = PptPt(13)
    qp1.font.bold = True
    qp1.font.color.rgb = PptRGBColor(6, 95, 70)
    qp1.alignment = PP_ALIGN.CENTER
    qp2 = q_tf8.add_paragraph()
    qp2.text = "手機直接掃碼互動"
    qp2.font.size = PptPt(10)
    qp2.font.color.rgb = PptRGBColor(100, 116, 139)
    qp2.alignment = PP_ALIGN.CENTER

    pptx_path = os.path.join(WORK_DIR, "花蓮生活探索地圖_開發歷程與心得簡報_4比3.pptx")
    prs.save(pptx_path)
    print("Saved 4:3 PPTX:", pptx_path)
    return pptx_path

# ==========================================
# 2. 建立 4:3 HTML 簡報並轉成 4:3 PDF
# ==========================================
def create_43_pdf_slides():
    html_content = """<!DOCTYPE html>
<html lang="zh-TW">
<head>
<meta charset="UTF-8">
<title>花蓮生活探索地圖 - 開發歷程與心得簡報 (4:3)</title>
<style>
  @page {
    size: 10in 7.5in landscape;
    margin: 0;
  }
  * {
    box-sizing: border-box;
    margin: 0;
    padding: 0;
  }
  body {
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "PingFang TC", "Microsoft JhengHei", sans-serif;
    background: #0f172a;
    color: #334155;
  }
  .slide {
    width: 10in;
    height: 7.5in;
    margin: 0 auto;
    background: #ffffff;
    padding: 40px 50px;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
    position: relative;
    page-break-after: always;
  }
  .slide-title-bg {
    background: linear-gradient(135deg, #064e3b 0%, #065f46 50%, #047857 100%);
    color: #ffffff;
  }
  .header-box h2 {
    font-size: 24pt;
    font-weight: 800;
    color: #0f172a;
    display: flex;
    align-items: center;
    gap: 10px;
  }
  .header-box p {
    font-size: 11pt;
    color: #64748b;
    margin-top: 4px;
    border-bottom: 2px solid #e2e8f0;
    padding-bottom: 10px;
  }
  .content-grid-2 {
    display: grid;
    grid-template-columns: 1fr 1fr;
    gap: 24px;
    flex: 1;
    margin: 18px 0;
  }
  .content-grid-3 {
    display: grid;
    grid-template-columns: 1fr 1fr 1fr;
    gap: 16px;
    flex: 1;
    margin: 18px 0;
  }
  .card {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 12px;
    padding: 16px 20px;
    display: flex;
    flex-direction: column;
    gap: 8px;
  }
  .card h3 {
    font-size: 14pt;
    font-weight: 700;
    color: #0f172a;
  }
  .card p, .card li {
    font-size: 10.5pt;
    line-height: 1.55;
    color: #475569;
  }
  .card ul {
    padding-left: 18px;
  }
  .card-highlight {
    background: #ecfdf5;
    border-color: #a7f3d0;
  }
  .card-danger {
    background: #fef2f2;
    border-color: #fecaca;
  }
  .card-warning {
    background: #fff7ed;
    border-color: #fed7aa;
  }
  .card-purple {
    background: #f5f3ff;
    border-color: #ddd6fe;
  }
  .card-amber {
    background: #fffbeb;
    border-color: #fde68a;
  }
  .footer-bar {
    display: flex;
    justify-content: space-between;
    font-size: 9pt;
    color: #94a3b8;
    border-top: 1px solid #f1f5f9;
    padding-top: 8px;
  }
  .badge {
    display: inline-block;
    padding: 3px 10px;
    border-radius: 9999px;
    font-size: 9pt;
    font-weight: 700;
    background: #e2e8f0;
    color: #475569;
  }
</style>
</head>
<body>

  <!-- 第 1 頁：封面 -->
  <section class="slide slide-title-bg" style="justify-content:center; align-items:center;">
    <div style="display:flex; width:100%; align-items:center; justify-content:space-between; gap:40px;">
      <div style="flex:1;">
        <span style="display:inline-block; background:rgba(255,255,255,0.2); padding:5px 14px; border-radius:20px; font-size:12pt; margin-bottom:16px;">期中專題 · 開發歷程與人機協作全紀錄</span>
        <h1 style="font-size:38pt; font-weight:800; margin-bottom:12px; line-height:1.15;">花蓮生活探索地圖</h1>
        <h2 style="font-size:17pt; color:#a7f3d0; font-weight:400; margin-bottom:24px;">從零發想、踩坑排雷到最終落地實踐</h2>
        <div style="font-size:11.5pt; color:#ecfdf5; line-height:1.8; border-top:1px solid rgba(255,255,255,0.25); padding-top:16px;">
          <div>• <strong>主題分類</strong>：蔬食美饌 · 精選景點 · 質感咖啡館指南</div>
          <div>• <strong>技術架構</strong>：HTML5 / Vanilla JS / Leaflet / Esri 街道圖資</div>
          <div>• <strong>協同開發</strong>：Antigravity (Google DeepMind) 協作</div>
        </div>
      </div>
      <div style="background:#ffffff; padding:18px 22px; border-radius:16px; text-align:center; color:#0f172a; box-shadow:0 12px 28px rgba(0,0,0,0.3); width:200px;">
        <img src="qrcode.png" style="width:160px; height:160px; display:block; margin:0 auto 10px auto; border-radius:6px;">
        <div style="font-weight:700; font-size:12pt; color:#065f46;">📱 掃碼即刻體驗</div>
        <div style="font-size:8pt; color:#64748b; word-break:break-all; margin-top:2px;">411513004-coder.github.io/1028_-</div>
      </div>
    </div>
  </section>

  <!-- 第 2 頁：專案起點 -->
  <section class="slide">
    <div class="header-box">
      <h2>🚀 專案起點：為什麼做？如何出發？</h2>
      <p>Project Inception & Core Requirements · 兼顧課堂規範與真實生活需求</p>
    </div>
    <div class="content-grid-2">
      <div class="card card-highlight">
        <h3 style="color:#065f46;">📌 課堂要求與現實限制</h3>
        <ul>
          <li><strong>90 秒即開即評</strong>：展示必須在 90 秒內講完「打開什麼、做什麼、看到什麼」，系統不能有龐大相依環境或開機等待。</li>
          <li><strong>零金鑰資安紅線</strong>：程式碼、截圖與遠端倉庫嚴禁含有任何 API 金鑰，否則重扣。</li>
          <li><strong>零安裝零負擔</strong>：任一電腦雙擊 <code>index.html</code> 即可於瀏覽器秒開執行。</li>
        </ul>
      </div>
      <div class="card">
        <h3>💡 初衷發想：花蓮綠色生活指南</h3>
        <ul>
          <li><strong>解決蔬食痛點</strong>：傳統地圖將「純素」、「蛋奶素」與「葷素共食」混淆，素食者常需反覆電詢確認。</li>
          <li><strong>學生與旅客剛需</strong>：東華大學、慈濟生活圈缺乏整合性在地地圖，外地旅客常找不到靠譜推薦。</li>
          <li><strong>決定尋求 AI 協同</strong>：由 Antigravity 快速產出系統雛形與前端骨架，開啟人機協作之旅。</li>
        </ul>
      </div>
    </div>
    <div class="footer-bar">
      <span>花蓮生活探索地圖 · 4:3 簡報</span>
      <span class="badge">2 / 8 頁</span>
    </div>
  </section>

  <!-- 第 3 頁：遭遇困難一：AI 幻覺與店名自創 -->
  <section class="slide">
    <div class="header-box">
      <h2>🚨 遭遇困難一：AI 幻覺、幽靈店家與店名自創</h2>
      <p>Challenge 1: AI Hallucinations & Ghost Stores · 速度飛快，但資料暗藏危機</p>
    </div>
    <div class="content-grid-2">
      <div class="card card-danger">
        <h3 style="color:#b91c1c;">👻 幽靈店家：抓取已停業名單</h3>
        <p><strong>【現象】</strong>AI 剛開始蒐集資料速度極快，但在細查後發現：</p>
        <ul>
          <li>七星潭人氣老宅「<strong>龍宮咖啡</strong>」官方 IG 早已宣布結束實體營業！</li>
          <li>搜尋名單充斥已歇業或轉型的過期資料。</li>
        </ul>
        <p><strong>【排雷】</strong>人工覆核最新社群動態，剔除停業店家，替補新城老街正派經營之「練習曲書店咖啡」。</p>
      </div>
      <div class="card card-warning">
        <h3 style="color:#c2410c;">🏷️ 憑空捏造：幫現存店家創生新名字</h3>
        <p><strong>【現象】</strong>店家明明存在，AI 卻自己給了「平行時空新名稱」：</p>
        <ul>
          <li>網路上查無任何登記紀錄的「太魯閣旅人素食小館」。</li>
          <li>壽豐東華商圈將「阿宏素食」誤植為「宏銘素食」。</li>
        </ul>
        <p><strong>【排雷】</strong>透過在地生活經驗比對花蓮縣政府登錄資料，將新城名店更正為工藝師名店「織娘之家無菜單料理」。</p>
      </div>
    </div>
    <div class="footer-bar">
      <span>花蓮生活探索地圖 · 4:3 簡報</span>
      <span class="badge">3 / 8 頁</span>
    </div>
  </section>

  <!-- 第 4 頁：遭遇困難二：地圖底圖踩坑與排錯 -->
  <section class="slide">
    <div class="header-box">
      <h2>🗺️ 遭遇困難二：底圖破圖、浮水印與協議限制</h2>
      <p>Challenge 2: Basemap Watermark & Zoom Breakdown · 零金鑰圖資的實戰考驗</p>
    </div>
    <div class="content-grid-3">
      <div class="card">
        <h3>1. CARTO 浮水印限制</h3>
        <p><strong>【問題】</strong>在 local 端直接以 <code>file://</code> 開啟時，CARTO 圖資會自動覆蓋滿版「CARTO」水印，視覺受到嚴重干擾。</p>
        <p><strong>【原因】</strong>CARTO 對無 Referrer 或本機檔案協議有安全防護防盜鏈。</p>
      </div>
      <div class="card">
        <h3>2. NatGeo 放大極限破圖</h3>
        <p><strong>【問題】</strong>換用國家地理風圖資，但放大到市區街道（Zoom 16 以上）時，畫面全黑並提示「Map data not yet available」。</p>
        <p><strong>【原因】</strong>該底圖瓦片最大僅支援至 16 級縮放，無法滿足市區小吃定位。</p>
      </div>
      <div class="card card-highlight">
        <h3 style="color:#065f46;">✅ 最終最佳解：Esri 方案</h3>
        <p><strong>【解法】</strong>全面重構為 Esri 雙圖層：</p>
        <ul>
          <li><strong>主圖</strong>：Esri World Street Map（街道圖清晰無水印，支援高解析 Zoom 19 級）。</li>
          <li><strong>副圖</strong>：Esri World Imagery（高空衛星圖），100% 零商業金鑰、穩定度 100%。</li>
        </ul>
      </div>
    </div>
    <div class="footer-bar">
      <span>花蓮生活探索地圖 · 4:3 簡報</span>
      <span class="badge">4 / 8 頁</span>
    </div>
  </section>

  <!-- 第 5 頁：功能擴充與三大主題成形 -->
  <section class="slide">
    <div class="header-box">
      <h2>🌿 功能擴充：升級「花蓮生活探索地圖」</h2>
      <p>Scope Expansion: From Veggie Map to Comprehensive Living Guide · 50 處全實名核實</p>
    </div>
    <div class="content-grid-3">
      <div class="card card-highlight">
        <h3 style="color:#065f46;">🌱 蔬食美饌 (30 家)</h3>
        <ul>
          <li><strong>全素/純素 (16家)</strong>：蝸牛蔬食Buffet、植地佛陀碗、明園煎餃、宜蘭水煎包、東華阿宏素食、織娘之家。</li>
          <li><strong>蛋奶素 (4家)</strong>：Piatto+菜老宅義式、青田の蔬食、草緣齋披薩。</li>
          <li><strong>蔬食友善 (10家)</strong>：時光1939有機早午餐、櫻之田野野菜鍋、好佳小吃(獨立素食鍋具)。</li>
        </ul>
      </div>
      <div class="card card-purple">
        <h3 style="color:#6d28d9;">🏛️ 精選景點 (10 處)</h3>
        <ul>
          <li><strong>歷史人文</strong>：將軍府1936、松園別館(2025新團隊重啟獲百大亮點)、文創園區舊酒廠、東大門夜市、吉安慶修院。</li>
          <li><strong>自然山海</strong>：七星潭月牙灣、壽豐鯉魚潭、雲山水落羽松、秀林台泥DAKA園區、崇德瑩農場。</li>
        </ul>
      </div>
      <div class="card card-amber">
        <h3 style="color:#b45309;">☕ 質感咖啡館 (10 家)</h3>
        <ul>
          <li><strong>老宅與烘焙</strong>：浮室純白老宅、Giocare手沖陶藝、珈琲花乾燥花、KŌHI宅深夜榻榻米、Stroll散步咖啡肉桂捲、SENSES烘焙。</li>
          <li><strong>古蹟與山海</strong>：咖逼小売所(縣定古蹟)、森山舍、練習曲書店、海碉堡。</li>
        </ul>
      </div>
    </div>
    <div class="footer-bar">
      <span>花蓮生活探索地圖 · 4:3 簡報</span>
      <span class="badge">5 / 8 頁</span>
    </div>
  </section>

  <!-- 第 6 頁：製作細節與技術亮點 -->
  <section class="slide">
    <div class="header-box">
      <h2>⚙️ 製作細節：技術亮點與互動設計</h2>
      <p>Technical Implementation & UI/UX Highlights · 打造流暢的使用者體驗</p>
    </div>
    <div class="content-grid-2">
      <div class="card">
        <h3 style="color:#065f46;">✨ 雙層智慧導航系統</h3>
        <p>首頁提供四大主題大標籤（全部 / 蔬食 / 景點 / 咖啡廳），並具備自動感知機制：僅在瀏覽蔬食時動態展開全素、蛋奶素、友善子分類標籤，介面整潔乾淨。</p>
      </div>
      <div class="card">
        <h3 style="color:#065f46;">✨ 雙向卡片地圖聯動</h3>
        <p>點擊左側卡片，地圖平滑平移（FlyTo）並展開詳細彈窗；點擊地圖標記，清單自動平滑滾動並高亮該卡片，資訊直觀連貫。</p>
      </div>
      <div class="card">
        <h3 style="color:#065f46;">✨ 多維分色自訂標記</h3>
        <p>綠色 🌱 全素、黃色 🥚 蛋奶素、藍色 🥗 友善、紫色 🏛️ 景點、焦糖色 ☕ 咖啡館，地圖全景一眼區隔，辨識效率極高。</p>
      </div>
      <div class="card">
        <h3 style="color:#065f46;">✨ 行動端無縫切換與外連</h3>
        <p>支援手機自適應抽屜（一鍵切換清單與地圖）；整合 <code>tel:</code> 一鍵通話與 Google Maps GPS 經緯度路徑導航。</p>
      </div>
    </div>
    <div class="footer-bar">
      <span>花蓮生活探索地圖 · 4:3 簡報</span>
      <span class="badge">6 / 8 頁</span>
    </div>
  </section>

  <!-- 第 7 頁：倒數第二頁 · 我的心得 -->
  <section class="slide">
    <div class="header-box">
      <h2>💬 專案心得與人機協作反思</h2>
      <p>Reflection on AI-Assisted Development · 速度與真實性的深刻體悟</p>
    </div>
    <div style="margin:20px 0; background:#ecfdf5; border:2px solid #10b981; border-radius:14px; padding:24px 28px;">
      <div style="font-size:13pt; font-weight:700; color:#065f46; margin-bottom:10px;">「開發者的真實心聲」</div>
      <blockquote style="font-size:13.5pt; font-weight:600; line-height:1.75; color:#0f172a;">
        這次我請 Antigravity 幫我架設網站，架設的速度快，但是在修正資料正確性卻花了不少時間，他會找到已經停業的店，或者雖然店存在，但是幫它創造了新名字，還有地圖的問題等說不完的問題。<br><br>
        在這次製作網站的過程中，雖然 AI 確實可以取代部份的工作，但是資料校正跟最後成品測試還是需要有人來把關。
      </blockquote>
    </div>
    <div class="content-grid-2" style="margin:0;">
      <div class="card" style="padding:12px 16px;">
        <h3 style="font-size:11.5pt; color:#10b981;">⚡ AI 的優勢：架構與代碼產出神速</h3>
        <p style="font-size:9.5pt;">從無到有建立完整的 Leaflet 互動地圖、雙向連動、CSS 響應式佈局與 GitHub Pages 自動部署，省去數天寫代碼的摸索。</p>
      </div>
      <div class="card" style="padding:12px 16px;">
        <h3 style="font-size:11.5pt; color:#ea580c;">🛡️ 人類的價值：嚴格把關與真實校驗</h3>
        <p style="font-size:9.5pt;">AI 缺乏真實世界驗證能力，容易把舊資料當現狀或自我造詞。唯有依靠人類的地圖實測、電話覆核與邏輯把關，才能產出真正有溫度的實用作品。</p>
      </div>
    </div>
    <div class="footer-bar">
      <span>花蓮生活探索地圖 · 4:3 簡報</span>
      <span class="badge">7 / 8 頁</span>
    </div>
  </section>

  <!-- 第 8 頁：最後一頁 · 謝謝大家 -->
  <section class="slide slide-title-bg" style="justify-content:center; align-items:center;">
    <div style="display:flex; width:100%; align-items:center; justify-content:space-between; gap:40px;">
      <div style="flex:1;">
        <span style="display:inline-block; background:rgba(255,255,255,0.2); padding:5px 14px; border-radius:20px; font-size:13pt; margin-bottom:16px;">報告完畢 · 敬請指教</span>
        <h1 style="font-size:46pt; font-weight:800; margin-bottom:16px; line-height:1.15;">謝謝大家！</h1>
        <p style="font-size:13pt; color:#d1fae5; line-height:1.8; margin-bottom:18px;">
          歡迎現場掃描右側 QR Code 體驗，或在 GitHub 查看完整無金鑰開原始碼。
        </p>
        <div style="font-size:11pt; color:#ecfdf5; line-height:1.7;">
          <div>🔗 <strong>專案首頁</strong>：https://411513004-coder.github.io/1028_-/</div>
          <div>📦 <strong>開源倉庫</strong>：https://github.com/411513004-coder/1028_-</div>
        </div>
      </div>
      <div style="background:#ffffff; padding:20px 24px; border-radius:18px; text-align:center; color:#0f172a; box-shadow:0 12px 28px rgba(0,0,0,0.3); width:210px;">
        <img src="qrcode.png" style="width:160px; height:160px; display:block; margin:0 auto 10px auto; border-radius:6px;">
        <div style="font-weight:700; font-size:12pt; color:#065f46;">花蓮生活探索地圖</div>
        <div style="font-size:8.5pt; color:#64748b; margin-top:3px;">手機掃碼直接互動</div>
      </div>
    </div>
  </section>

</body>
</html>"""

    html_path = os.path.join(WORK_DIR, "_tmp_43_slides.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)

    pdf_path = os.path.join(WORK_DIR, "花蓮生活探索地圖_開發歷程與心得簡報_4比3.pdf")
    cmd = [
        EDGE_PATH,
        '--headless',
        '--disable-gpu',
        '--no-pdf-header-footer',
        f'--print-to-pdf={pdf_path}',
        html_path
    ]
    subprocess.run(cmd, capture_output=True)
    if os.path.exists(html_path):
        os.remove(html_path)
    print("Saved 4:3 PDF Slides:", pdf_path)
    return pdf_path

# ==========================================
# 3. 建立完整歷程與心得 Word 文件 (.docx)
# ==========================================
def create_retro_doc():
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run("《花蓮生活探索地圖》專題開發歷程、技術細節與人機協作心得報告")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(6, 78, 59)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(12)

    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("專案名稱", "花蓮生活探索地圖 (Hualien Living Map) —— 蔬食 · 景點 · 咖啡館指南"),
        ("協同工具", "Antigravity (Google DeepMind) 協作開發"),
        ("專案規格", "100% 零商業金鑰 (No API Key) · 零安裝雙擊即跑 · 50 處全實名核實"),
        ("開源倉庫", "https://github.com/411513004-coder/1028_-"),
        ("線上展示", "https://411513004-coder.github.io/1028_-/")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(1.5)
        c2.width = Inches(5.2)
        set_cell_background(c1, "F0FDF4")
        set_cell_background(c2, "F8FAFC")
        set_cell_margins(c1, 80, 80, 120, 120)
        set_cell_margins(c2, 80, 80, 120, 120)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(k)
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(6, 95, 70)
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(v)
        r2.font.size = Pt(10)
        if "http" in v:
            r2.font.color.rgb = RGBColor(2, 132, 199)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    def add_sec_h(num_str, title_str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(f"{num_str}、 {title_str}")
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(6, 78, 59)
        return h

    def add_sub_h(title_str):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(title_str)
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
        return h

    # 1. 專案緣起與需求背景
    add_sec_h("一", "專案緣起與需求背景")
    p = doc.add_paragraph()
    p.add_run("本專案始於大專院校期中專題的實作挑戰。課堂規範十分嚴謹明確：每位同學須於現場進行「90 秒實機演示」，說明「打開什麼、做什麼、看到什麼」，且評分特別強調「專案必須保證能在展示當天順暢運行，不強制安裝龐大軟體或笨重伺服器環境」，最關鍵的一條紅線則是「程式碼與截圖中嚴禁出現任何商業 API 金鑰（如 Google Maps API Key），杜絕資安外洩與計費超額風險」。").font.size = Pt(10.5)

    p2 = doc.add_paragraph()
    p2.add_run("在主題發想上，鑑於花蓮在地居民、大專院校師生（如東華大學、慈濟大學）以及自由行觀光客常面臨「蔬食分類標籤混雜（純素、蛋奶素與一般葷食餐廳混淆）」以及「缺乏整合性慢活旅遊地圖」的痛點，我決定藉由現代 Web 前端技術，打造一個兼具蔬食美饌、必訪景點與自習老宅咖啡館的輕量化互動地圖應用 ——《花蓮生活探索地圖》。").font.size = Pt(10.5)

    # 2. 開發歷程與遭遇困難排雷
    add_sec_h("二", "開發歷程與遭遇困難排雷（踩坑與解決實錄）")
    p = doc.add_paragraph()
    p.add_run("在本次專題開發中，我邀請了 Antigravity 作為協同開發 AI。雖然 AI 在代碼生成的速度上表現驚人，但在真實落地的過程中，卻遭遇了許多意料之外的困難與挫折：").font.size = Pt(10.5)

    add_sub_h("1. 遭遇困難一：AI 幻覺與幽靈店家（已停業）")
    p = doc.add_paragraph()
    p.add_run("AI 剛開始快速蒐集並產出店家清單，乍看格式非常工整完備，然而當我進行實地與官方社群查驗時，卻發現了嚴重的「幽靈店家」地雷：\n").font.size = Pt(10.5)
    bullets1 = [
        ("已停業名店被當成現況：", "例如七星潭知名老宅名店「龍宮咖啡」，AI 一開始仍將其收錄，但透過官方 Instagram 查核發現該店早已結束實體營運。"),
        ("解決策略：", "我立即要求 AI 建立嚴格查核機制，全面剔除已停業店家，並替補新城老街正派經營且充滿教育關懷的人氣空間「練習曲書店咖啡」，確保所有清單 100% 現正營運。")
    ]
    for b_title, b_desc in bullets1:
        bp = doc.add_paragraph(style='List Bullet')
        r_t = bp.add_run(b_title)
        r_t.font.bold = True
        r_t.font.size = Pt(10)
        r_d = bp.add_run(b_desc)
        r_d.font.size = Pt(10)

    add_sub_h("2. 遭遇困難二：店名憑空捏造與名稱誤植（AI 幻覺）")
    p = doc.add_paragraph()
    p.add_run("除了過期停業店家外，更棘手的是 AI 會在資料庫中產生「AI 幻覺（Hallucination）」：\n").font.size = Pt(10.5)
    bullets2 = [
        ("憑空發明不存在的店名：", "清單中一度出現「太魯閣旅人素食小館」，但我上網與花蓮商業登記系統反覆查找，花蓮根本沒有這家店。"),
        ("店家存在但自創新名字：", "在壽豐鄉東華大學商圈，AI 將知名的「阿宏素食」誤寫為「宏銘素食」。"),
        ("解決策略：", "我透過自身的在地生活經驗與實地資料比對，揪出這些虛構與混淆的名稱，將新城名店校正為太魯閣工藝師打造之「織娘之家無菜單料理」，並精準更正阿宏素食、好佳小吃等真實口碑名店。")
    ]
    for b_title, b_desc in bullets2:
        bp = doc.add_paragraph(style='List Bullet')
        r_t = bp.add_run(b_title)
        r_t.font.bold = True
        r_t.font.size = Pt(10)
        r_d = bp.add_run(b_desc)
        r_d.font.size = Pt(10)

    add_sub_h("3. 遭遇困難三：地圖圖資的浮水印與放大破圖")
    p = doc.add_paragraph()
    p.add_run("為遵守「零金鑰」規範，在地圖選型上經歷了多輪波折：\n").font.size = Pt(10.5)
    bullets3 = [
        ("CARTO Voyager 浮水印干擾：", "一開始採用 CARTO 圖資，但在本機透過 file:// 協議直接點擊 HTML 執行時，CARTO 會偵測無 Referrer 並打上滿版浮水印，嚴重破壞簡報展示質感。"),
        ("NatGeo 國家地理圖資拉近黑屏：", "隨後嘗試切換為 NatGeo 圖資，但在拉近到市區尋找精準巷弄時（Zoom 16 以上），畫面直接破圖黑屏並顯示「Map data not yet available」。"),
        ("最終最佳解：", "全面換裝開源 Leaflet 搭配 Esri World Street Map 街道圖資與高空空照圖，不僅完全免金鑰、無浮水印，且能順暢放大到 19 級精準定位巷弄。")
    ]
    for b_title, b_desc in bullets3:
        bp = doc.add_paragraph(style='List Bullet')
        r_t = bp.add_run(b_title)
        r_t.font.bold = True
        r_t.font.size = Pt(10)
        r_d = bp.add_run(b_desc)
        r_d.font.size = Pt(10)

    # 3. 系統核心架構與製作細節
    add_sec_h("三", "系統核心架構與製作細節")
    p = doc.add_paragraph()
    p.add_run("經過反覆調整後的最終系統具備以下技術亮點：").font.size = Pt(10.5)

    techs = [
        ("極致輕量純原生前端：", "HTML5 + CSS3 + Vanilla JavaScript，首頁秒級載入，支援桌機雙欄與手機端單欄自適應抽屜。"),
        ("三大主題與 50 處全核實據點：", "涵蓋 30 家蔬食（純素 16、蛋奶素 4、友善 10）、10 大精選景點（文化古蹟與自然山海）與 10 家質感咖啡館（老宅、烘焙、海景）。"),
        ("雙層智慧篩選與全文即時搜尋：", "切換主題時動態連動子分類標籤；搜尋框支援店名、地址與招牌特色（如「肉桂捲」、「將軍府」、「水煎包」）即時檢索。"),
        ("雙向連動與一鍵行動外連：", "卡片點擊自動 FlyTo 定位並展開彈窗；Pin 點擊自動滾動清單；支援一鍵撥號與 Google Maps 路線規劃。"),
        ("GitHub Pages 自動部署與專屬 QR Code：", "建置專屬線上站台，並生成專屬 QR Code 圖檔，讓台下觀眾能直接拿出手機掃描互動。")
    ]
    for t_title, t_desc in techs:
        tp = doc.add_paragraph(style='List Bullet')
        r_t = tp.add_run(t_title)
        r_t.font.bold = True
        r_t.font.size = Pt(10)
        r_d = tp.add_run(t_desc)
        r_d.font.size = Pt(10)

    # 4. 專題心得與人機協作反思（完整引用並延伸）
    add_sec_h("四", "專題心得與人機協作反思（開發者心得）")
    
    # 建立引言框
    quote_table = doc.add_table(rows=1, cols=1)
    quote_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = quote_table.rows[0].cells[0]
    c.width = Inches(6.5)
    set_cell_background(c, "ECFDF5")
    set_cell_margins(c, 140, 140, 180, 180)
    qp = c.paragraphs[0]
    qr_t = qp.add_run("【開發者的真實體悟】\n")
    qr_t.font.bold = True
    qr_t.font.size = Pt(11)
    qr_t.font.color.rgb = RGBColor(6, 95, 70)
    
    qr_body = qp.add_run("「這次我請 Antigravity 幫我架設網站，架設的速度快，但是在修正資料正確性卻花了不少時間，他會找到已經停業的店，或者雖然店存在，但是幫它創造了新名字，還有地圖的問題等說不完的問題。\n\n在這次製作網站的過程中，雖然 AI 確實可以取代部份的工作，但是資料校正跟最後成品測試還是需要有人來把關。」")
    qr_body.font.size = Pt(10.5)
    qr_body.font.bold = True
    qr_body.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    p_reflect = doc.add_paragraph()
    p_reflect.add_run("這段歷程給予我最深刻的啟示是：現階段的生成式 AI 就像一位速度極快、技術博學的初階工程師，能夠在幾分鐘內完成複雜的介面切版、互動邏輯與底圖串接；然而，AI 對「真實物理世界」的感知是滯後且脆弱的，缺乏實地求證的能力。\n\n真正讓一個軟體從「Demo 半成品」昇華為「真正可信賴的產品」，關鍵在於身為專案主導者的人類。從使用者痛點的定義、資料真實性的實地覆核、地圖破圖時的邏輯推理，到最後成品的把關與打磨，人類的批判性思維與審美判斷，依然是不可取代的核心價值。").font.size = Pt(10.5)

    docx_path = os.path.join(WORK_DIR, "花蓮生活探索地圖_專題開發歷程與心得報告.docx")
    doc.save(docx_path)
    print("Saved Retro DOCX:", docx_path)
    return docx_path

# ==========================================
# 4. 建立心得報告 PDF 檔 (.pdf)
# ==========================================
def create_retro_pdf():
    html_content = """<!DOCTYPE html>
<html lang="zh-TW">
<head>
<meta charset="UTF-8">
<title>《花蓮生活探索地圖》專題開發歷程與心得報告</title>
<style>
  @page {
    size: A4;
    margin: 20mm 18mm 20mm 18mm;
  }
  body {
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "PingFang TC", "Microsoft JhengHei", sans-serif;
    color: #1e293b;
    line-height: 1.65;
    font-size: 10.5pt;
  }
  h1 {
    color: #064e3b;
    text-align: center;
    font-size: 19pt;
    margin-bottom: 6px;
    letter-spacing: -0.5px;
  }
  .meta-box {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    padding: 12px 18px;
    margin: 14px 0 20px 0;
    display: grid;
    grid-template-columns: auto 1fr;
    gap: 6px 14px;
    font-size: 9.5pt;
  }
  .meta-label {
    font-weight: bold;
    color: #065f46;
  }
  h2 {
    color: #064e3b;
    border-bottom: 2px solid #a7f3d0;
    padding-bottom: 4px;
    margin-top: 20px;
    margin-bottom: 8px;
    font-size: 13pt;
  }
  h3 {
    color: #0f172a;
    font-size: 11pt;
    margin-top: 12px;
    margin-bottom: 4px;
  }
  ul {
    margin-top: 4px;
    margin-bottom: 8px;
    padding-left: 20px;
  }
  li {
    margin-bottom: 5px;
  }
  .quote-box {
    background: #ecfdf5;
    border: 2px solid #10b981;
    border-radius: 8px;
    padding: 14px 18px;
    margin: 14px 0;
  }
  .quote-box .q-title {
    font-weight: bold;
    color: #065f46;
    margin-bottom: 6px;
    font-size: 10.5pt;
  }
  .quote-box blockquote {
    font-size: 11pt;
    font-weight: 600;
    line-height: 1.7;
    color: #0f172a;
  }
  .qr-float {
    float: right;
    width: 110px;
    text-align: center;
    margin: 0 0 10px 15px;
    padding: 8px;
    border: 1px solid #e2e8f0;
    border-radius: 8px;
    background: #ffffff;
  }
  .qr-float img {
    width: 90px;
    height: 90px;
    display: block;
    margin: 0 auto 4px auto;
  }
  .qr-float span {
    font-size: 7.5pt;
    color: #64748b;
  }
</style>
</head>
<body>

<div class="qr-float">
  <img src="qrcode.png" alt="QR Code">
  <span>📱 掃碼線上體驗</span>
</div>

<h1>《花蓮生活探索地圖》專題開發歷程與心得報告</h1>

<div class="meta-box">
  <div class="meta-label">專案名稱：</div><div>花蓮生活探索地圖 (Hualien Living Map)</div>
  <div class="meta-label">協同工具：</div><div>Antigravity (Google DeepMind) 協作開發</div>
  <div class="meta-label">專案規格：</div><div>100% 零商業金鑰 (No API Key) · 零安裝雙擊即跑 · 50 處全實名核實</div>
  <div class="meta-label">開源倉庫：</div><div>https://github.com/411513004-coder/1028_-</div>
  <div class="meta-label">線上展示：</div><div>https://411513004-coder.github.io/1028_-/</div>
</div>

<h2>一、 專案緣起與需求背景</h2>
<p>本專案始於大專院校期中專題的實作挑戰。課堂規範十分嚴謹明確：每位同學須於現場進行「90 秒實機演示」，說明「打開什麼、做什麼、看到什麼」，且評分特別強調「專案必須保證能在展示當天順暢運行，不強制安裝龐大軟體或笨重伺服器環境」，最關鍵的一條紅線則是「程式碼與截圖中嚴禁出現任何商業 API 金鑰（如 Google Maps API Key），杜絕資安外洩與計費超額風險」。</p>
<p>在主題發想上，鑑於花蓮在地居民、大專院校師生（如東華大學、慈濟大學）以及自由行觀光客常面臨「蔬食分類標籤混雜（純素、蛋奶素與一般葷食餐廳混淆）」以及「缺乏整合性慢活旅遊地圖」的痛點，我決定藉由現代 Web 前端技術，打造一個兼具蔬食美饌、必訪景點與自習老宅咖啡館的輕量化互動地圖應用 ——《花蓮生活探索地圖》。</p>

<h2>二、 開發歷程與遭遇困難排雷（踩坑與解決實錄）</h2>
<p>在本次專題開發中，我邀請了 Antigravity 作為協同開發 AI。雖然 AI 在代碼生成的速度上表現驚人，但在真實落地的過程中，卻遭遇了許多意料之外的困難與挫折：</p>

<h3>1. 遭遇困難一：AI 幻覺與幽靈店家（已停業）</h3>
<ul>
  <li><strong>已停業名店被當成現況</strong>：例如七星潭知名老宅名店「龍宮咖啡」，AI 一開始仍將其收錄，但透過官方 Instagram 查核發現該店早已結束實體營運。</li>
  <li><strong>解決策略</strong>：我立即要求 AI 建立嚴格查核機制，全面剔除已停業店家，並替補新城老街正派經營且充滿教育關懷的人氣空間「練習曲書店咖啡」，確保所有清單 100% 現正營運。</li>
</ul>

<h3>2. 遭遇困難二：店名憑空捏造與名稱誤植（AI 幻覺）</h3>
<ul>
  <li><strong>憑空發明不存在的店名</strong>：清單中一度出現「太魯閣旅人素食小館」，但我上網與花蓮商業登記系統反覆查找，花蓮根本沒有這家店。</li>
  <li><strong>店家存在但自創新名字</strong>：在壽豐鄉東華大學商圈，AI 將知名的「阿宏素食」誤寫為「宏銘素食」。</li>
  <li><strong>解決策略</strong>：我透過自身的在地生活經驗與實地資料比對，揪出這些虛構與混淆的名稱，將新城名店校正為太魯閣工藝師打造之「織娘之家無菜單料理」，並精準更正阿宏素食、好佳小吃等真實口碑名店。</li>
</ul>

<h3>3. 遭遇困難三：地圖圖資的浮水印與放大破圖</h3>
<ul>
  <li><strong>CARTO Voyager 浮水印干擾</strong>：一開始採用 CARTO 圖資，但在本機透過 <code>file://</code> 協議直接點擊 HTML 執行時，CARTO 會偵測無 Referrer 並打上滿版浮水印，嚴重破壞簡報展示質感。</li>
  <li><strong>NatGeo 國家地理圖資拉近黑屏</strong>：隨後嘗試切換為 NatGeo 圖資，但在拉近到市區尋找精準巷弄時（Zoom 16 以上），畫面直接破圖黑屏並顯示「Map data not yet available」。</li>
  <li><strong>最終最佳解</strong>：全面換裝開源 Leaflet 搭配 Esri World Street Map 街道圖資與高空空照圖，不僅完全免金鑰、無浮水印，且能順暢放大到 19 級精準定位巷弄。</li>
</ul>

<h2>三、 系統核心架構與製作細節</h2>
<ul>
  <li><strong>極致輕量純原生前端</strong>：HTML5 + CSS3 + Vanilla JavaScript，首頁秒級載入，支援桌機雙欄與手機端單欄自適應抽屜。</li>
  <li><strong>三大主題與 50 處全核實據點</strong>：涵蓋 30 家蔬食（純素 16、蛋奶素 4、友善 10）、10 大精選景點（文化古蹟與自然山海）與 10 家質感咖啡館（老宅、烘焙、海景）。</li>
  <li><strong>雙層智慧篩選與全文即時搜尋</strong>：切換主題時動態連動子分類標籤；搜尋框支援店名、地址與招牌特色（如「肉桂捲」、「將軍府」、「水煎包」）即時檢索。</li>
  <li><strong>雙向連動與一鍵行動外連</strong>：卡片點擊自動 FlyTo 定位並展開彈窗；Pin 點擊自動滾動清單；支援一鍵撥號與 Google Maps 路線規劃。</li>
  <li><strong>GitHub Pages 自動部署與專屬 QR Code</strong>：建置專屬線上站台，並生成專屬 QR Code 圖檔，讓台下觀眾能直接拿出手機掃描互動。</li>
</ul>

<h2>四、 專題心得與人機協作反思（開發者心得）</h2>

<div class="quote-box">
  <div class="q-title">【開發者的真實心聲】</div>
  <blockquote>
    「這次我請 Antigravity 幫我架設網站，架設的速度快，但是在修正資料正確性卻花了不少時間，他會找到已經停業的店，或者雖然店存在，但是幫它創造了新名字，還有地圖的問題等說不完的問題。<br><br>
    在這次製作網站的過程中，雖然 AI 確實可以取代部份的工作，但是資料校正跟最後成品測試還是需要有人來把關。」
  </blockquote>
</div>

<p>這段歷程給予我最深刻的啟示是：現階段的生成式 AI 就像一位速度極快、技術博學的初階工程師，能夠在幾分鐘內完成複雜的介面切版、互動邏輯與底圖串接；然而，AI 對「真實物理世界」的感知是滯後且脆弱的，缺乏實地求證的能力。</p>
<p>真正讓一個軟體從「Demo 半成品」昇華為「真正可信賴的產品」，關鍵在於身為專案主導者的人類。從使用者痛點的定義、資料真實性的實地覆核、地圖破圖時的邏輯推理，到最後成品的把關與打磨，人類的批判性思維與審美判斷，依然是不可取代的核心價值。</p>

</body>
</html>"""

    html_path = os.path.join(WORK_DIR, "_tmp_retro.html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)

    pdf_path = os.path.join(WORK_DIR, "花蓮生活探索地圖_專題開發歷程與心得報告.pdf")
    cmd = [
        EDGE_PATH,
        '--headless',
        '--disable-gpu',
        '--no-pdf-header-footer',
        f'--print-to-pdf={pdf_path}',
        html_path
    ]
    subprocess.run(cmd, capture_output=True)
    if os.path.exists(html_path):
        os.remove(html_path)
    print("Saved Retro PDF:", pdf_path)
    return pdf_path

if __name__ == "__main__":
    create_43_pptx()
    create_43_pdf_slides()
    create_retro_doc()
    create_retro_pdf()
    print("All 4:3 deliverables generated successfully!")
